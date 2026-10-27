"""
Generador de Informe Diario de Inspección y Memoria Fotográfica en Excel (.xlsx).
Utiliza openpyxl y Pillow para insertar tablas formateadas, metadatos y fotografías en alta calidad.
"""
import os
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional

import openpyxl
from openpyxl.drawing.image import Image as OpenpyxlImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from PIL import Image as PILImage


def _create_temp_resized_image(src_path: str, max_w_px: int = 340, max_h_px: int = 240) -> Optional[str]:
    """Redimensiona proporcionalmente la imagen para que encaje limpia en la celda de Excel."""
    if not os.path.exists(src_path):
        return None

    try:
        with PILImage.open(src_path) as img:
            orig_w, orig_h = img.size
            if orig_w <= 0 or orig_h <= 0:
                return None

            ratio = min(max_w_px / orig_w, max_h_px / orig_h)
            new_w = max(1, int(orig_w * ratio))
            new_h = max(1, int(orig_h * ratio))

            resized = img.resize((new_w, new_h), PILImage.Resampling.LANCZOS)

            temp_dir = Path(src_path).parent / "temp_excel"
            temp_dir.mkdir(parents=True, exist_ok=True)
            temp_path = temp_dir / f"xl_{Path(src_path).name}"
            resized.save(temp_path, quality=85)
            return str(temp_path)
    except Exception as exc:
        print(f"Error optimizando imagen para Excel: {exc}")
        return None


def generate_inspection_excel(
    report_data: Dict[str, Any],
    photos: List[Dict[str, Any]],
    output_excel_path: str,
) -> str:
    """
    Genera un libro de Excel (.xlsx) con la Ficha Técnica, Actividades y Memoria Fotográfica.
    """
    os.makedirs(os.path.dirname(os.path.abspath(output_excel_path)), exist_ok=True)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Ficha de Inspección"
    ws.views.sheetView[0].showGridLines = True

    # Paleta de colores de ingeniería
    COLOR_PRIMARY_NAVY = "1E3A8A"     # Azul Marino Encabezado
    COLOR_SECONDARY_AMBER = "B45309"  # Ámbar Actividades
    COLOR_TEAL = "0F766E"             # Verde Fotos
    COLOR_SLATE_HEADER = "334155"     # Gris Encabezado Ficha
    COLOR_BG_LIGHT = "F8FAFC"         # Fondo suave celdas
    COLOR_BORDER = "CBD5E1"           # Borde sutil

    font_title = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    font_subtitle = Font(name="Calibri", size=10, bold=True, color="78350F")
    font_sec_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_field_label = Font(name="Calibri", size=10, bold=True, color="1E293B")
    font_field_value = Font(name="Calibri", size=10, color="0F172A")
    font_photo_header = Font(name="Calibri", size=9, bold=True, color="1E293B")
    font_photo_caption = Font(name="Calibri", size=8.5, italic=True, color="475569")

    thin_border_side = Side(style="thin", color=COLOR_BORDER)
    full_border = Border(
        left=thin_border_side,
        right=thin_border_side,
        top=thin_border_side,
        bottom=thin_border_side,
    )

    fill_title = PatternFill(start_color=COLOR_PRIMARY_NAVY, end_color=COLOR_PRIMARY_NAVY, fill_type="solid")
    fill_subtitle = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
    fill_sec_ficha = PatternFill(start_color=COLOR_SLATE_HEADER, end_color=COLOR_SLATE_HEADER, fill_type="solid")
    fill_sec_act = PatternFill(start_color=COLOR_SECONDARY_AMBER, end_color=COLOR_SECONDARY_AMBER, fill_type="solid")
    fill_sec_foto = PatternFill(start_color=COLOR_TEAL, end_color=COLOR_TEAL, fill_type="solid")
    fill_row_label = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    fill_card_header = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")

    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)

    # Anchos de columna
    col_widths = {
        "A": 3,
        "B": 24,
        "C": 28,
        "D": 4,
        "E": 24,
        "F": 28,
        "G": 3,
    }
    for col_letter, width in col_widths.items():
        ws.column_dimensions[col_letter].width = width

    # 1. ENCABEZADO
    ws.merge_cells("B2:F2")
    ws["B2"] = "INFORME DIARIO DE INSPECCIÓN TÉCNICA DE OBRA"
    ws["B2"].font = font_title
    ws["B2"].fill = fill_title
    ws["B2"].alignment = align_center
    ws.row_dimensions[2].height = 28

    ws.merge_cells("B3:F3")
    ws["B3"] = "CONTROL DE AVANCE FÍSICO Y MEMORIA FOTOGRÁFICA"
    ws["B3"].font = font_subtitle
    ws["B3"].fill = fill_subtitle
    ws["B3"].alignment = align_center
    ws.row_dimensions[3].height = 20

    current_row = 5

    # 2. SECCIÓN: FICHA TÉCNICA
    ws.merge_cells(f"B{current_row}:F{current_row}")
    cell_sec1 = ws[f"B{current_row}"]
    cell_sec1.value = "FICHA TÉCNICA DEL PROYECTO"
    cell_sec1.font = font_sec_header
    cell_sec1.fill = fill_sec_ficha
    cell_sec1.alignment = align_center
    ws.row_dimensions[current_row].height = 22
    current_row += 1

    # Fila Obra (Ocupa ancho completo)
    ws[f"B{current_row}"] = "OBRA / PROYECTO:"
    ws[f"B{current_row}"].font = font_field_label
    ws[f"B{current_row}"].fill = fill_row_label
    ws[f"B{current_row}"].border = full_border

    ws.merge_cells(f"C{current_row}:F{current_row}")
    cell_obra = ws[f"C{current_row}"]
    cell_obra.value = report_data.get("obra", "SIN ESPECIFICAR").upper()
    cell_obra.font = font_field_value
    cell_obra.alignment = align_left
    cell_obra.border = full_border
    ws.row_dimensions[current_row].height = 36
    current_row += 1

    # Fila Fecha y Empresa Contratista
    ws[f"B{current_row}"] = "FECHA DEL REPORTE:"
    ws[f"B{current_row}"].font = font_field_label
    ws[f"B{current_row}"].fill = fill_row_label
    ws[f"B{current_row}"].border = full_border

    ws[f"C{current_row}"] = report_data.get("fecha", datetime.now().strftime("%d/%m/%Y"))
    ws[f"C{current_row}"].font = font_field_value
    ws[f"C{current_row}"].alignment = align_left
    ws[f"C{current_row}"].border = full_border

    ws[f"E{current_row}"] = "EMPRESA CONTRATISTA:"
    ws[f"E{current_row}"].font = font_field_label
    ws[f"E{current_row}"].fill = fill_row_label
    ws[f"E{current_row}"].border = full_border

    ws[f"F{current_row}"] = report_data.get("empresa", "No indicada")
    ws[f"F{current_row}"].font = font_field_value
    ws[f"F{current_row}"].alignment = align_left
    ws[f"F{current_row}"].border = full_border
    ws.row_dimensions[current_row].height = 20
    current_row += 1

    # Fila Residente e Inspectores
    inspectores_list = report_data.get("ing_inspectores", [])
    inspectores_str = ", ".join(inspectores_list) if inspectores_list else "No indicados"

    ws[f"B{current_row}"] = "ING. RESIDENTE:"
    ws[f"B{current_row}"].font = font_field_label
    ws[f"B{current_row}"].fill = fill_row_label
    ws[f"B{current_row}"].border = full_border

    ws[f"C{current_row}"] = report_data.get("ing_residente", "No indicado")
    ws[f"C{current_row}"].font = font_field_value
    ws[f"C{current_row}"].alignment = align_left
    ws[f"C{current_row}"].border = full_border

    ws[f"E{current_row}"] = "ING. INSPECTORES:"
    ws[f"E{current_row}"].font = font_field_label
    ws[f"E{current_row}"].fill = fill_row_label
    ws[f"E{current_row}"].border = full_border

    ws[f"F{current_row}"] = inspectores_str
    ws[f"F{current_row}"].font = font_field_value
    ws[f"F{current_row}"].alignment = align_left
    ws[f"F{current_row}"].border = full_border
    ws.row_dimensions[current_row].height = 20
    current_row += 1

    # Fila Caporal y Fuerza Laboral
    obreros_text = report_data.get("obreros", "No especificado")
    caporal_obra_num = report_data.get("caporal_obra", "")
    personal_str = f"{obreros_text}"
    if caporal_obra_num:
        personal_str += f" • Caporal: {caporal_obra_num}"

    ws[f"B{current_row}"] = "CAPORAL DE OBRA:"
    ws[f"B{current_row}"].font = font_field_label
    ws[f"B{current_row}"].fill = fill_row_label
    ws[f"B{current_row}"].border = full_border

    ws[f"C{current_row}"] = report_data.get("caporal", "No indicado")
    ws[f"C{current_row}"].font = font_field_value
    ws[f"C{current_row}"].alignment = align_left
    ws[f"C{current_row}"].border = full_border

    ws[f"E{current_row}"] = "FUERZA LABORAL:"
    ws[f"E{current_row}"].font = font_field_label
    ws[f"E{current_row}"].fill = fill_row_label
    ws[f"E{current_row}"].border = full_border

    ws[f"F{current_row}"] = personal_str
    ws[f"F{current_row}"].font = font_field_value
    ws[f"F{current_row}"].alignment = align_left
    ws[f"F{current_row}"].border = full_border
    ws.row_dimensions[current_row].height = 20
    current_row += 2

    # 3. SECCIÓN: ACTIVIDADES EN EJECUCIÓN
    ws.merge_cells(f"B{current_row}:F{current_row}")
    cell_sec2 = ws[f"B{current_row}"]
    cell_sec2.value = "ACTIVIDADES EJECUTADAS Y EN DESARROLLO"
    cell_sec2.font = font_sec_header
    cell_sec2.fill = fill_sec_act
    cell_sec2.alignment = align_center
    ws.row_dimensions[current_row].height = 22
    current_row += 1

    actividades = report_data.get("actividades", [])
    if actividades:
        for idx, act in enumerate(actividades, 1):
            ws.merge_cells(f"B{current_row}:F{current_row}")
            cell_act = ws[f"B{current_row}"]
            cell_act.value = f"  {idx}.  {act}"
            cell_act.font = font_field_value
            cell_act.alignment = align_left
            cell_act.border = full_border
            ws.row_dimensions[current_row].height = 22
            current_row += 1
    else:
        ws.merge_cells(f"B{current_row}:F{current_row}")
        cell_none = ws[f"B{current_row}"]
        cell_none.value = "  Sin actividades reportadas para la fecha."
        cell_none.font = Font(name="Calibri", size=9.5, italic=True)
        cell_none.alignment = align_left
        cell_none.border = full_border
        ws.row_dimensions[current_row].height = 20
        current_row += 1

    current_row += 1

    # 4. SECCIÓN: MEMORIA FOTOGRÁFICA
    if photos:
        ws.merge_cells(f"B{current_row}:F{current_row}")
        cell_sec3 = ws[f"B{current_row}"]
        cell_sec3.value = f"MEMORIA FOTOGRÁFICA DE LA JORNADA ({len(photos)} FOTOGRAFÍAS)"
        cell_sec3.font = font_sec_header
        cell_sec3.fill = fill_sec_foto
        cell_sec3.alignment = align_center
        ws.row_dimensions[current_row].height = 22
        current_row += 2

        # Organizar fotos en 2 columnas: Col B y Col E
        for i in range(0, len(photos), 2):
            batch = photos[i : i + 2]

            # Fila de encabezado de fotos
            header_row = current_row
            ws.merge_cells(f"B{header_row}:C{header_row}")
            ws[f"B{header_row}"] = f"FOTOGRAFÍA N° {i + 1}"
            ws[f"B{header_row}"].font = font_photo_header
            ws[f"B{header_row}"].fill = fill_card_header
            ws[f"B{header_row}"].alignment = align_center
            ws[f"B{header_row}"].border = full_border

            if len(batch) > 1:
                ws.merge_cells(f"E{header_row}:F{header_row}")
                ws[f"E{header_row}"] = f"FOTOGRAFÍA N° {i + 2}"
                ws[f"E{header_row}"].font = font_photo_header
                ws[f"E{header_row}"].fill = fill_card_header
                ws[f"E{header_row}"].alignment = align_center
                ws[f"E{header_row}"].border = full_border

            ws.row_dimensions[header_row].height = 18

            # Fila del contenedor de la imagen
            img_row = header_row + 1
            ws.row_dimensions[img_row].height = 160

            # Insertar Foto 1
            p1 = batch[0]
            opt_p1 = _create_temp_resized_image(p1.get("path", ""), max_w_px=340, max_h_px=200)
            if opt_p1 and os.path.exists(opt_p1):
                img1 = OpenpyxlImage(opt_p1)
                ws.add_image(img1, f"B{img_row}")

            # Insertar Foto 2 si existe
            if len(batch) > 1:
                p2 = batch[1]
                opt_p2 = _create_temp_resized_image(p2.get("path", ""), max_w_px=340, max_h_px=200)
                if opt_p2 and os.path.exists(opt_p2):
                    img2 = OpenpyxlImage(opt_p2)
                    ws.add_image(img2, f"E{img_row}")

            # Fila de leyendas / pie de fotos
            caption_row = img_row + 1
            cap1 = p1.get("caption", "").strip() or f"Actividad en fecha {report_data.get('fecha', '')}"
            ws.merge_cells(f"B{caption_row}:C{caption_row}")
            ws[f"B{caption_row}"] = cap1
            ws[f"B{caption_row}"].font = font_photo_caption
            ws[f"B{caption_row}"].alignment = align_center
            ws[f"B{caption_row}"].border = full_border

            if len(batch) > 1:
                cap2 = batch[1].get("caption", "").strip() or f"Actividad en fecha {report_data.get('fecha', '')}"
                ws.merge_cells(f"E{caption_row}:F{caption_row}")
                ws[f"E{caption_row}"] = cap2
                ws[f"E{caption_row}"].font = font_photo_caption
                ws[f"E{caption_row}"].alignment = align_center
                ws[f"E{caption_row}"].border = full_border

            ws.row_dimensions[caption_row].height = 24
            current_row = caption_row + 2

    # 5. CUADRO DE FIRMAS AL FINAL
    current_row += 1
    ws.merge_cells(f"B{current_row}:C{current_row}")
    ws[f"B{current_row}"] = "POR LA INSPECCIÓN DE OBRA"
    ws[f"B{current_row}"].font = font_field_label
    ws[f"B{current_row}"].alignment = align_center
    ws[f"B{current_row}"].fill = fill_row_label
    ws[f"B{current_row}"].border = full_border

    ws.merge_cells(f"E{current_row}:F{current_row}")
    ws[f"E{current_row}"] = "POR LA EMPRESA CONTRATISTA"
    ws[f"E{current_row}"].font = font_field_label
    ws[f"E{current_row}"].alignment = align_center
    ws[f"E{current_row}"].fill = fill_row_label
    ws[f"E{current_row}"].border = full_border
    ws.row_dimensions[current_row].height = 20

    current_row += 1
    # Espacio para firmar
    ws.merge_cells(f"B{current_row}:C{current_row + 2}")
    ws.merge_cells(f"E{current_row}:F{current_row + 2}")
    for r in range(current_row, current_row + 3):
        ws.row_dimensions[r].height = 18

    current_row += 3
    # Nombres bajo firma
    ws.merge_cells(f"B{current_row}:C{current_row}")
    ws[f"B{current_row}"] = f"{inspectores_str}\nFirma y Sello / C.I.V."
    ws[f"B{current_row}"].font = font_photo_caption
    ws[f"B{current_row}"].alignment = align_center

    ws.merge_cells(f"E{current_row}:F{current_row}")
    ws[f"E{current_row}"] = f"{report_data.get('ing_residente', '')} (Residente)\nFirma y Sello / C.I.V."
    ws[f"E{current_row}"].font = font_photo_caption
    ws[f"E{current_row}"].alignment = align_center
    ws.row_dimensions[current_row].height = 26

    wb.save(output_excel_path)
    return output_excel_path
