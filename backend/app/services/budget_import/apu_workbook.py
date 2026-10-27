"""
Extractor especializado de presupuestos a partir de Libros de Excel con Hojas de APUs.
Analiza libros con múltiples hojas donde cada hoja contiene un Análisis de Precio Unitario (APU)
con su encabezado (Obra, Contratante, Código, Descripción, Rendimiento, Unidad, Cantidad) y sus
secciones completas de MATERIALES, EQUIPOS y MANO DE OBRA.
"""
import io
import logging
import re
import unicodedata
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import openpyxl
from openpyxl.worksheet.worksheet import Worksheet

logger = logging.getLogger(__name__)

MAX_APU_SHEETS = 300

_UNIT_MAP = {
    "m2": "m2", "m²": "m2", "mts2": "m2", "mt2": "m2", "metros cuadrados": "m2",
    "m3": "m3", "m³": "m3", "mts3": "m3", "mt3": "m3", "metros cubicos": "m3",
    "ml": "m", "m": "m", "mts": "m", "mt": "m", "metro": "m", "metros": "m",
    "und": "und", "un": "und", "ud": "und", "u": "und", "unidad": "und", "pza": "pza",
    "kg": "kg", "kgs": "kg", "kgf": "kgf", "kilogramo": "kg", "kilos": "kg",
    "ton": "ton", "tonf": "ton", "tonelada": "ton", "sco": "sco", "saco": "sco",
    "gal": "gal", "gln": "gal", "lt": "lt", "l": "lt", "litro": "lt",
    "dia": "día", "día": "día", "dias": "día", "hr": "hr", "hora": "hr",
    "gl": "gl", "global": "gl", "pto": "pto", "punto": "pto",
    "m3-km": "m3-km", "m3 km": "m3-km", "ton-km": "ton-km", "viaje": "viaje",
}


def _normalize_unit(value: Any) -> str:
    """Normaliza unidades comunes de construcción."""
    if not value:
        return "und"
    clean = _clean_str(value).lower()
    clean_no_accent = _strip_accents(clean)
    if clean_no_accent in _UNIT_MAP:
        return _UNIT_MAP[clean_no_accent]
    if clean in _UNIT_MAP:
        return _UNIT_MAP[clean]
    return clean[:15]


def _strip_accents(value: str) -> str:
    """Elimina tildes y diacríticos para comparaciones robustas."""
    if not value:
        return ""
    normalized = unicodedata.normalize("NFKD", str(value))
    return "".join(ch for ch in normalized if not unicodedata.combining(ch))


def _clean_str(value: Any) -> str:
    """Limpia cadenas eliminando espacios extras y valores nulos."""
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%d/%m/%Y")
    text = str(value).strip()
    if text.lower() in {"none", "nan", "nat"}:
        return ""
    return re.sub(r"\s+", " ", text)


def _to_float(value: Any, default: float = 0.0) -> float:
    """Convierte de forma segura celdas a flotante, soportando formatos venezolanos/españoles."""
    if value is None:
        return default
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace("Bs.", "").replace("Bs", "").replace("$", "").replace("%", "")
    text = text.replace(" ", "")
    if not text:
        return default
    # Manejar formatos como 1.234,56 -> 1234.56 o 1234.56
    if "," in text and "." in text:
        if text.rfind(",") > text.rfind("."):
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", "")
    elif "," in text:
        text = text.replace(",", ".")
    try:
        return float(text)
    except (ValueError, TypeError):
        return default


def is_apu_workbook(data: bytes) -> bool:
    """
    Detecta de forma no destructiva si un archivo Excel contiene hojas estructuradas de APU.
    Busca firmas como 'ANALISIS DE PRECIO UNITARIO' o bloques simultáneos de MATERIALES y MANO DE OBRA.
    """
    if not data:
        return False
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
        sheet_names = wb.sheetnames[:10]
        for name in sheet_names:
            ws = wb[name]
            found_mat = False
            found_mo = False
            for row in ws.iter_rows(max_row=35, values_only=True):
                for cell in row:
                    if cell is None:
                        continue
                    clean_upper = _strip_accents(str(cell)).upper()
                    if "ANALISIS DE PRECIO" in clean_upper:
                        return True
                    if "MATERIALES" in clean_upper and len(clean_upper) < 25:
                        found_mat = True
                    if "MANO DE OBRA" in clean_upper and len(clean_upper) < 25:
                        found_mo = True
            if found_mat and found_mo:
                return True
        return False
    except Exception as exc:
        logger.error("Error al verificar si el Excel es libro de APUs: %s", exc, exc_info=True)
        return False


def _is_apu_sheet(ws: Worksheet) -> bool:
    """Verifica si una hoja individual de Excel corresponde a un APU."""
    max_r = min(ws.max_row or 50, 40)
    found_mat = False
    found_mo = False
    for r in range(1, max_r + 1):
        for c in range(1, min(ws.max_column or 10, 10) + 1):
            val = ws.cell(r, c).value
            if val is None:
                continue
            clean_upper = _strip_accents(str(val)).upper()
            if "ANALISIS DE PRECIO" in clean_upper:
                return True
            if "MATERIALES" in clean_upper and len(clean_upper) < 25:
                found_mat = True
            if "MANO DE OBRA" in clean_upper and len(clean_upper) < 25:
                found_mo = True
    return found_mat and found_mo


def _parse_single_apu_sheet(ws: Worksheet, sheet_name: str, sheet_index: int) -> Optional[Dict[str, Any]]:
    """
    Extrae la información completa de una hoja de APU (metadatos, materiales, equipos y mano de obra).
    """
    max_r = ws.max_row or 60
    max_c = min(ws.max_column or 12, 12)
    cells = [[ws.cell(r, c).value for c in range(1, max_c + 1)] for r in range(1, max_r + 1)]

    obra: Optional[str] = None
    contratante: Optional[str] = None
    fecha: Optional[str] = None
    part_no: Optional[str] = None
    description: Optional[str] = None
    performance: float = 1.0
    code: Optional[str] = None
    unit: str = "und"
    quantity: float = 1.0

    fcas: Optional[float] = None
    admin: Optional[float] = None
    profit: Optional[float] = None
    iva: Optional[float] = None
    currency: str = "USD"

    materials: List[Dict[str, Any]] = []
    equipments: List[Dict[str, Any]] = []
    labors: List[Dict[str, Any]] = []

    mode: Optional[str] = None

    for row in cells:
        row_str = " ".join([str(v) for v in row if v is not None])
        row_clean = _strip_accents(row_str).upper()

        # Detección de moneda
        if any(term in row_clean for term in ("BS.F", "BS.", "BOLIVAR", "BOLIVARES")):
            currency = "VES"

        # Escaneo de metadatos de cabecera
        for c_idx, val in enumerate(row):
            if val is None:
                continue
            sval = _clean_str(val)
            supper = _strip_accents(sval).upper()

            if "OBRA:" in supper and not obra:
                if ":" in sval and len(sval.split(":", 1)[1].strip()) > 1:
                    obra = sval.split(":", 1)[1].strip()
                elif c_idx + 1 < len(row) and row[c_idx + 1] is not None:
                    obra = _clean_str(row[c_idx + 1])

            if any(k in supper for k in ("CONTRATANTE:", "CLIENTE:")) and not contratante:
                if ":" in sval and len(sval.split(":", 1)[1].strip()) > 1:
                    contratante = sval.split(":", 1)[1].strip()
                elif c_idx + 1 < len(row) and row[c_idx + 1] is not None:
                    contratante = _clean_str(row[c_idx + 1])

            if "FECHA:" in supper and not fecha:
                if ":" in sval and len(sval.split(":", 1)[1].strip()) > 1:
                    fecha = sval.split(":", 1)[1].strip()
                elif c_idx + 1 < len(row) and row[c_idx + 1] is not None:
                    fecha = _clean_str(row[c_idx + 1])

            if any(k in supper for k in ("PART. NO.:", "PARTIDA NO", "PARTIDA:", "PART. N")) and not part_no:
                if ":" in sval and len(sval.split(":", 1)[1].strip()) > 0:
                    part_no = sval.split(":", 1)[1].strip()
                elif c_idx + 1 < len(row) and row[c_idx + 1] is not None:
                    part_no = _clean_str(row[c_idx + 1])

            if "DESCRIPCI" in supper and not description:
                if ":" in sval and len(sval.split(":", 1)[1].strip()) > 2:
                    description = sval.split(":", 1)[1].strip()
                elif c_idx + 1 < len(row) and row[c_idx + 1] is not None:
                    description = _clean_str(row[c_idx + 1])

            if "RENDIMIENTO:" in supper:
                if ":" in sval and len(sval.split(":", 1)[1].strip()) > 0:
                    val_parsed = _to_float(sval.split(":", 1)[1])
                    if val_parsed > 0:
                        performance = val_parsed
                elif c_idx + 1 < len(row) and row[c_idx + 1] is not None:
                    val_parsed = _to_float(row[c_idx + 1])
                    if val_parsed > 0:
                        performance = val_parsed

            if any(k in supper for k in ("CODIGO:", "CODIGO DE PARTIDA", "COD:")) and not code:
                if ":" in sval and len(sval.split(":", 1)[1].strip()) > 0:
                    code = sval.split(":", 1)[1].strip()
                elif c_idx + 1 < len(row) and row[c_idx + 1] is not None:
                    code = _clean_str(row[c_idx + 1])

            if "UNIDAD:" in supper:
                if ":" in sval and len(sval.split(":", 1)[1].strip()) > 0:
                    unit = sval.split(":", 1)[1].strip()
                elif c_idx + 1 < len(row) and row[c_idx + 1] is not None:
                    unit = _clean_str(row[c_idx + 1])

            if "CANTIDAD:" in supper:
                if ":" in sval and len(sval.split(":", 1)[1].strip()) > 0:
                    val_parsed = _to_float(sval.split(":", 1)[1])
                    if val_parsed >= 0:
                        quantity = val_parsed
                elif c_idx + 1 < len(row) and row[c_idx + 1] is not None:
                    val_parsed = _to_float(row[c_idx + 1])
                    if val_parsed >= 0:
                        quantity = val_parsed

            # Recargos e indirectos
            if "PRESTACIONES SOCIALES" in supper and fcas is None:
                for cand in row:
                    cand_f = _to_float(cand, default=-1.0)
                    if 50.0 <= cand_f <= 1500.0:
                        fcas = cand_f
                        break

            if "ADMINISTRACI" in supper and admin is None:
                for cand in row:
                    cand_f = _to_float(cand, default=-1.0)
                    if 1.0 <= cand_f <= 50.0:
                        admin = cand_f
                        break

            if "UTILIDAD" in supper and profit is None:
                for cand in row:
                    cand_f = _to_float(cand, default=-1.0)
                    if 1.0 <= cand_f <= 50.0:
                        profit = cand_f
                        break

            if ("IMPUESTO" in supper or "I.V.A." in supper) and iva is None:
                for cand in row:
                    cand_f = _to_float(cand, default=-1.0)
                    if 0.0 <= cand_f <= 35.0:
                        iva = cand_f
                        break

        # Cambio de sección
        if "MATERIALES" in row_clean and len(row_clean.strip()) < 25:
            mode = "MATERIALES"
            continue
        elif "EQUIPOS" in row_clean and len(row_clean.strip()) < 25:
            mode = "EQUIPOS"
            continue
        elif "MANO DE OBRA" in row_clean and len(row_clean.strip()) < 30:
            mode = "MANO DE OBRA"
            continue

        if any(term in row_clean for term in ("TOTAL MATERIALES", "TOTAL EQUIPOS", "SUBTOTAL MANO DE OBRA", "TOTAL GENERAL MANO DE OBRA", "COSTO DIRECTO")):
            mode = None
            continue

        # Extracción de insumos según sección
        if mode == "MATERIALES":
            if any(term in row_clean for term in ("NO.", "DESCRIPCI", "UND.", "PRECIO")):
                continue
            col_code = _clean_str(row[0]) if len(row) > 0 and row[0] is not None else ""
            desc = _clean_str(row[2]) if len(row) > 2 and row[2] is not None else ""
            if not desc and len(row) > 1 and isinstance(row[1], str) and len(str(row[1])) > 3:
                desc = _clean_str(row[1])

            if desc and len(desc) >= 2 and not "TOTAL" in desc.upper():
                und_val = _clean_str(row[3]) if len(row) > 3 and row[3] is not None else "und"
                cant_val = _to_float(row[4] if len(row) > 4 else 0.0, default=0.0)
                desp_val = _to_float(row[5] if len(row) > 5 else 0.0, default=0.0)
                precio_val = _to_float(row[6] if len(row) > 6 else 0.0, default=0.0)
                materials.append({
                    "codigo": col_code or f"MAT-{len(materials) + 1:03d}",
                    "descripcion": desc,
                    "unidad": _normalize_unit(und_val) or "und",
                    "cantidad": cant_val,
                    "desperdicio": desp_val,
                    "precio_unitario": precio_val,
                })

        elif mode == "EQUIPOS":
            if any(term in row_clean for term in ("NO.", "DESCRIPCI", "COP/DEP", "PRECIO")):
                continue
            col_code = _clean_str(row[0]) if len(row) > 0 and row[0] is not None else ""
            desc = _clean_str(row[2]) if len(row) > 2 and row[2] is not None else ""
            if not desc and len(row) > 1 and isinstance(row[1], str) and len(str(row[1])) > 3:
                desc = _clean_str(row[1])

            if desc and len(desc) >= 2 and not "TOTAL" in desc.upper():
                cant_val = _to_float(row[4] if len(row) > 4 else 1.0, default=1.0)
                cop_val = _to_float(row[5] if len(row) > 5 else 1.0, default=1.0)
                precio_val = _to_float(row[6] if len(row) > 6 else 0.0, default=0.0)
                equipments.append({
                    "codigo": col_code or f"EQ-{len(equipments) + 1:03d}",
                    "descripcion": desc,
                    "unidad": "día",
                    "cantidad": cant_val,
                    "depreciacion": cop_val if cop_val > 0 else 1.0,
                    "precio_unitario": precio_val,
                })

        elif mode == "MANO DE OBRA":
            if any(term in row_clean for term in ("NO.", "DESCRIPCI", "JORNAL", "BONO")):
                continue
            col_code = _clean_str(row[0]) if len(row) > 0 and row[0] is not None else ""
            desc = _clean_str(row[2]) if len(row) > 2 and row[2] is not None else ""
            if not desc and len(row) > 1 and isinstance(row[1], str) and len(str(row[1])) > 3:
                desc = _clean_str(row[1])

            if desc and len(desc) >= 2 and not any(t in desc.upper() for t in ("TOTAL", "SUBTOTAL")):
                cant_val = _to_float(row[3] if len(row) > 3 else 1.0, default=1.0)
                jornal_val = _to_float(row[4] if len(row) > 4 else 0.0, default=0.0)
                bono_val = _to_float(row[5] if len(row) > 5 else 0.0, default=0.0)
                labors.append({
                    "codigo": col_code or f"MO-{len(labors) + 1:03d}",
                    "descripcion": desc,
                    "cantidad": cant_val,
                    "jornal": jornal_val,
                    "bono": bono_val,
                })

    if not description and not materials and not equipments and not labors:
        return None

    final_code = code or part_no or f"APU-{sheet_index:03d}"
    final_desc = description or f"Partida {final_code}"

    return {
        "obra": obra,
        "contratante": contratante,
        "fecha": fecha,
        "part_no": part_no,
        "code": final_code,
        "description": final_desc,
        "performance": max(0.01, performance),
        "unit": _normalize_unit(unit) or "und",
        "quantity": max(0.0, quantity),
        "currency": currency,
        "fcas": fcas,
        "admin": admin,
        "profit": profit,
        "iva": iva,
        "materials": materials,
        "equipments": equipments,
        "labors": labors,
    }


def extract_from_apu_workbook(data: bytes, filename: str) -> Dict[str, Any]:
    """
    Procesa un libro de Excel completo que contiene hojas de APUs y construye la estructura
    de vista previa para el creador de presupuestos.
    """
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    items: List[Dict[str, Any]] = []

    global_obra: Optional[str] = None
    global_contratante: Optional[str] = None
    global_currency: str = "USD"
    global_fcas: Optional[float] = None
    global_admin: Optional[float] = None
    global_profit: Optional[float] = None
    global_iva: Optional[float] = None

    warnings: List[str] = []

    sheet_names = wb.sheetnames[:MAX_APU_SHEETS]
    for idx, sheet_name in enumerate(sheet_names, start=1):
        ws = wb[sheet_name]
        if not _is_apu_sheet(ws):
            continue

        apu = _parse_single_apu_sheet(ws, sheet_name, idx)
        if not apu:
            continue

        if not global_obra and apu.get("obra"):
            global_obra = apu["obra"]
        if not global_contratante and apu.get("contratante"):
            global_contratante = apu["contratante"]
        if apu.get("currency") == "VES":
            global_currency = "VES"
        if global_fcas is None and apu.get("fcas"):
            global_fcas = apu["fcas"]
        if global_admin is None and apu.get("admin"):
            global_admin = apu["admin"]
        if global_profit is None and apu.get("profit"):
            global_profit = apu["profit"]
        if global_iva is None and apu.get("iva"):
            global_iva = apu["iva"]

        item_row = len(items) + 1
        item_code = str(apu["code"]).strip()
        item_desc = str(apu["description"]).strip().upper()
        item_unit = str(apu["unit"]).strip()
        item_qty = float(apu["quantity"])
        item_perf = float(apu["performance"])

        m_count = len(apu["materials"])
        e_count = len(apu["equipments"])
        l_count = len(apu["labors"])

        candidate = {
            "cod_par": item_code,
            "cov_par": item_code,
            "description": item_desc,
            "unit": item_unit,
            "score": 1.0,
            "match_type": "excel_apu",
        }

        items.append({
            "row": item_row,
            "code": item_code,
            "description": item_desc,
            "unit": item_unit,
            "quantity": item_qty,
            "performance": item_perf,
            "calculation": f"Hoja '{sheet_name}' • APU completo ({m_count} Mat, {e_count} Eq, {l_count} MO)",
            "confidence": "1.0",
            "warnings": [],
            "is_chapter": False,
            "candidates": [candidate],
            "selected_cod_par": item_code,
            "has_covenin_code_match": True,
            "apu_source": "excel_apu",
            "materials": apu["materials"],
            "equipments": apu["equipments"],
            "labors": apu["labors"],
        })

    if not items:
        raise ValueError("No se encontraron hojas con estructura válida de APU en el libro de Excel.")

    budget_name = global_obra or Path(filename or "").stem.replace("_", " ").title() or "Presupuesto APU Excel"

    header = {
        "budget_name": budget_name[:250],
        "project_name": global_obra[:250] if global_obra else None,
        "client_name": global_contratante[:250] if global_contratante else None,
        "currency": global_currency,
        "fcas_percent": global_fcas,
        "admin_percent": global_admin,
        "profit_percent": global_profit,
        "iva_percent": global_iva,
    }

    stats = {
        "total_items": len(items),
        "matched": len(items),
        "pending_ai": 0,
        "chapters": 0,
    }

    warnings.append(
        f"Se extrajeron exitosamente {len(items)} hojas de APUs completas con sus análisis de materiales, equipos y mano de obra."
    )

    return {
        "filename": filename,
        "source_type": "excel_apu",
        "header": header,
        "items": items,
        "stats": stats,
        "warnings": warnings,
    }
