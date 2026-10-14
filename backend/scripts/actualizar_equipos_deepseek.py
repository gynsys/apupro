import argparse
import json
import logging
import os
import sys
import time
from datetime import datetime
from typing import Any, Dict, List, Optional, Tuple

from dotenv import load_dotenv
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
import pandas as pd
import requests
from sqlalchemy import or_
from sqlalchemy.orm import Session

# Permitir ejecucion directa agregando la raiz del backend al sys.path
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
BACKEND_DIR = os.path.abspath(os.path.join(SCRIPT_DIR, ".."))
if BACKEND_DIR not in sys.path:
    sys.path.insert(0, BACKEND_DIR)

# Cargar variables de entorno del backend
load_dotenv(os.path.join(BACKEND_DIR, ".env"))

from app.crud.llm import decrypt_api_key
from app.db.models.costbase import CostEquipment
from app.db.models.llm_provider import LLMProvider
from app.db.session import SessionLocal

logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s")
logger = logging.getLogger("actualizar_equipos_deepseek")


def get_deepseek_credentials(db: Session) -> Tuple[str, str, str]:
    """
    Obtiene y desencripta las credenciales de DeepSeek desde la base de datos.
    Retorna (api_key, base_url, model_name).
    """
    if db is None:
        raise ValueError("La sesion de base de datos no puede ser nula.")

    provider = (
        db.query(LLMProvider)
        .filter(LLMProvider.provider_key == "deepseek", LLMProvider.is_active == True)
        .first()
    )
    if not provider:
        raise ValueError("No se encontro un proveedor activo para 'deepseek' en la tabla llm_providers.")

    api_key = decrypt_api_key(provider.api_key_enc)
    if not api_key:
        raise ValueError("La clave API de DeepSeek no pudo ser desencriptada o esta vacia.")

    base_url = "https://api.deepseek.com"
    if provider.extra_params:
        extra = provider.extra_params
        if isinstance(extra, str):
            try:
                extra = json.loads(extra)
            except Exception:
                extra = {}
        if isinstance(extra, dict) and extra.get("base_url"):
            base_url = extra["base_url"].rstrip("/")

    model_name = provider.model_name or "deepseek-chat"
    return api_key, base_url, model_name


def fetch_equipments_to_update(
    db: Session,
    limit: Optional[int] = None,
    search_filter: Optional[str] = None,
    deprec_filter: Optional[str] = None,
    only_zero_price: bool = False,
) -> List[Dict[str, Any]]:
    """
    Consulta los equipos a auditar y actualizar segun los filtros especificados.
    """
    query = db.query(CostEquipment)

    if only_zero_price:
        query = query.filter(or_(CostEquipment.precio.is_(None), CostEquipment.precio == 0.0))

    if search_filter:
        term = f"%{search_filter.strip()}%"
        query = query.filter(
            or_(
                CostEquipment.CodEqu.ilike(term),
                CostEquipment.ref_code.ilike(term),
                CostEquipment.Descri.ilike(term),
            )
        )

    if deprec_filter:
        clean = deprec_filter.strip().lower()
        if clean in ("equal_1", "1", "1.0", "alquiler"):
            query = query.filter(or_(CostEquipment.deprec_factor >= 0.99, CostEquipment.deprec_factor.is_(None)))
        elif clean in ("less_1", "<1", "propio"):
            query = query.filter(CostEquipment.deprec_factor < 0.99, CostEquipment.deprec_factor > 0)

    query = query.order_by(CostEquipment.ref_code.asc(), CostEquipment.CodEqu.asc())
    if limit and limit > 0:
        query = query.limit(limit)

    items: List[Dict[str, Any]] = []
    for eq in query.all():
        precio = float(eq.precio or 0.0)
        deprec = float(eq.deprec_factor if eq.deprec_factor is not None else 1.0)
        cosdia = float(eq.CosDia or round(precio * deprec, 4))

        items.append({
            "referencia": eq.ref_code or eq.CodEqu,
            "codigo_interno": eq.CodEqu,
            "descripcion": eq.Descri or "",
            "precio_actual": precio,
            "deprec_actual": deprec,
            "cosdia_actual": cosdia,
        })

    return items


def call_deepseek_equipment_batch(
    api_key: str,
    base_url: str,
    model_name: str,
    batch_items: List[Dict[str, Any]],
    max_retries: int = 3,
) -> List[Dict[str, Any]]:
    """
    Envia un lote de equipos a DeepSeek para clasificacion de modalidad y estimacion de precios de mercado.
    """
    if not batch_items:
        return []

    system_prompt = (
        "Eres un Ingeniero Mecánico y Director de Maquinaria y Equipos en Construcción y Obras Civiles.\n"
        "Tu misión es auditar y actualizar el costo de equipos y herramientas para Análisis de Precios Unitarios (APU) en USD.\n\n"
        "REGLAS OBLIGATORIAS DE INGENIERÍA:\n"
        "1. DISTINCIÓN DE MODALIDAD (PROPIO vs ALQUILER):\n"
        "   - HERRAMIENTAS MENORES / MANUALES (palas, picos, carretillas, mandarrias, barras, kits de albañilería, "
        "andamios simples, herramientas eléctricas menores portátiles): Son EQUIPO PROPIO AMORTIZABLE. "
        "Debes estimar el PRECIO DE COMPRA NUEVO de reposición de activo (USD) y asignar un factor de depreciación diaria "
        "real (desgaste del 1.5% al 2.5% diario, ej: 0.02 o 0.025; para andamios 0.008 diario). NUNCA dejes depreciación = 1.0 para herramientas menores.\n"
        "   - MAQUINARIA PESADA / MEDIANA (retroexcavadoras, camiones volteo, motoniveladoras, grúas, rodillos vibratorios, "
        "plantas eléctricas grandes): Si se maneja por TARIFA DE ALQUILER DIARIO, coloca modalidad 'ALQUILER', "
        "precio_nuevo = tarifa diaria de alquiler (USD/día) y factor_deprec = 1.0.\n"
        "2. COSTO DIARIO: En el JSON calcula siempre 'costo_diario' = round(precio_nuevo * factor_deprec, 2).\n"
        "3. FORMATO: Responde ÚNICAMENTE con un JSON con la clave 'items', donde cada elemento tenga:\n"
        "   'referencia': str,\n"
        "   'modalidad': 'PROPIO' o 'ALQUILER',\n"
        "   'precio_nuevo': float (valor compra nuevo si es PROPIO, o tarifa diaria si es ALQUILER),\n"
        "   'factor_deprec': float (entre 0.005 y 0.05 si es PROPIO; 1.0 si es ALQUILER),\n"
        "   'costo_diario': float,\n"
        "   'criterio': str (especificación, marca de referencia o desglose breve)\n"
    )

    clean_batch_for_prompt = [
        {
            "referencia": item["referencia"],
            "descripcion": item["descripcion"],
            "precio_actual": item["precio_actual"],
            "deprec_actual": item["deprec_actual"],
            "cosdia_actual": item["cosdia_actual"],
        }
        for item in batch_items
    ]

    user_prompt = (
        f"Analiza los siguientes {len(clean_batch_for_prompt)} equipos y herramientas de construcción. "
        f"Determina su modalidad correcta, precio de mercado nuevo en USD y depreciación diaria:\n\n"
        f"{json.dumps(clean_batch_for_prompt, ensure_ascii=False, indent=2)}\n\n"
        f"Devuelve un JSON válido con la estructura:\n"
        f'{{"items": [{{"referencia": "...", "modalidad": "PROPIO", "precio_nuevo": 0.0, "factor_deprec": 0.02, "costo_diario": 0.0, "criterio": "..."}}]}}'
    )

    payload = {
        "model": model_name,
        "messages": [
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": user_prompt},
        ],
        "temperature": 0.2,
        "max_tokens": 4096,
        "response_format": {"type": "json_object"},
    }

    headers = {
        "Authorization": f"Bearer {api_key}",
        "Content-Type": "application/json",
    }

    endpoint = f"{base_url}/chat/completions"

    for attempt in range(1, max_retries + 1):
        try:
            resp = requests.post(endpoint, headers=headers, json=payload, timeout=75)
            if resp.status_code == 429:
                wait_time = attempt * 4
                logger.warning(f"Rate limit de DeepSeek. Esperando {wait_time}s (intento {attempt}/{max_retries})...")
                time.sleep(wait_time)
                continue

            resp.raise_for_status()
            data = resp.json()
            content = data["choices"][0]["message"]["content"]
            parsed = json.loads(content)
            llm_items = parsed.get("items", [])

            ai_map = {str(it.get("referencia", "")).strip().upper(): it for it in llm_items if isinstance(it, dict)}

            enriched_results: List[Dict[str, Any]] = []
            for item in batch_items:
                ref_key = str(item["referencia"]).strip().upper()
                ai_data = ai_map.get(ref_key, {})

                modalidad = str(ai_data.get("modalidad") or "ALQUILER").upper()
                precio_nuevo = float(ai_data.get("precio_nuevo") or item["precio_actual"] or 0.0)
                factor_deprec = float(ai_data.get("factor_deprec") or item["deprec_actual"] or 1.0)

                # Salvaguarda: si es propio, el factor no puede ser 1.0
                if modalidad == "PROPIO" and factor_deprec >= 0.99:
                    factor_deprec = 0.02

                costo_diario_nuevo = round(precio_nuevo * factor_deprec, 2)
                cosdia_actual = float(item["cosdia_actual"] or 0.0)

                if cosdia_actual > 0:
                    var_cosdia_pct = round(((costo_diario_nuevo - cosdia_actual) / cosdia_actual) * 100.0, 2)
                else:
                    var_cosdia_pct = 100.0 if costo_diario_nuevo > 0 else 0.0

                criterio = str(ai_data.get("criterio") or "Estimacion de mercado DeepSeek V3")

                enriched_results.append({
                    "referencia": item["referencia"],
                    "codigo_interno": item["codigo_interno"],
                    "descripcion": item["descripcion"],
                    "modalidad": modalidad,
                    "precio_actual": item["precio_actual"],
                    "deprec_actual": item["deprec_actual"],
                    "cosdia_actual": cosdia_actual,
                    "precio_nuevo": round(precio_nuevo, 2),
                    "deprec_nuevo": round(factor_deprec, 4),
                    "cosdia_nuevo": costo_diario_nuevo,
                    "diferencia_cosdia": round(costo_diario_nuevo - cosdia_actual, 2),
                    "variacion_cosdia_pct": var_cosdia_pct,
                    "criterio": criterio,
                })

            return enriched_results

        except Exception as e:
            logger.error(f"Error llamando a DeepSeek en intento {attempt}/{max_retries}: {e}", exc_info=True)
            if attempt == max_retries:
                fallback_results = []
                for item in batch_items:
                    fallback_results.append({
                        "referencia": item["referencia"],
                        "codigo_interno": item["codigo_interno"],
                        "descripcion": item["descripcion"],
                        "modalidad": "ALQUILER" if item["deprec_actual"] >= 0.99 else "PROPIO",
                        "precio_actual": item["precio_actual"],
                        "deprec_actual": item["deprec_actual"],
                        "cosdia_actual": item["cosdia_actual"],
                        "precio_nuevo": item["precio_actual"],
                        "deprec_nuevo": item["deprec_actual"],
                        "cosdia_nuevo": item["cosdia_actual"],
                        "diferencia_cosdia": 0.0,
                        "variacion_cosdia_pct": 0.0,
                        "criterio": f"Fallo API DeepSeek: {e}",
                    })
                return fallback_results
            time.sleep(attempt * 2)

    return []


def export_equipments_to_excel(
    results: List[Dict[str, Any]],
    output_path: str,
) -> str:
    """
    Genera un archivo Excel profesional con los datos de auditoria y actualizacion de equipos.
    """
    if not results:
        raise ValueError("No hay resultados para exportar.")

    df_data = []
    for r in results:
        var = abs(r["variacion_cosdia_pct"])
        if r["cosdia_actual"] == 0:
            estado = "Precio Inicial Fijado"
        elif var > 50:
            estado = "Revisar (Variación > 50%)"
        elif var > 25:
            estado = "Alerta (Variación Moderada)"
        else:
            estado = "Aprobado Normal"

        df_data.append({
            "Referencia": r["referencia"],
            "Descripción": r["descripcion"],
            "Modalidad": r["modalidad"],
            "Precio Actual ($)": r["precio_actual"],
            "Deprec Actual": r["deprec_actual"],
            "Costo Diario Actual ($)": r["cosdia_actual"],
            "Precio Nuevo ($)": r["precio_nuevo"],
            "Deprec Nuevo": r["deprec_nuevo"],
            "Costo Diario Nuevo ($)": r["cosdia_nuevo"],
            "Diferencia Día ($)": r["diferencia_cosdia"],
            "Variación Día (%)": r["variacion_cosdia_pct"] / 100.0,
            "Criterio / Mercado": r["criterio"],
            "Estado Sugerido": estado,
            "Código Sistema": r["codigo_interno"],
        })

    df = pd.DataFrame(df_data)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Auditoría Equipos"

    # Titulo principal
    ws.merge_cells("A1:N1")
    title_cell = ws["A1"]
    title_cell.value = "COSTBASE PLATFORM — AUDITORÍA Y ACTUALIZACIÓN DE EQUIPOS Y MAQUINARIA"
    title_cell.font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    title_cell.fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 35

    # Subtitulo
    ws.merge_cells("A2:N2")
    sub_cell = ws["A2"]
    sub_cell.value = (
        f"Auditoría con DeepSeek V3 API | Fecha: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')} | "
        f"Total Equipos Auditados: {len(results)}"
    )
    sub_cell.font = Font(name="Calibri", size=10, italic=True, color="475569")
    sub_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 20

    headers = list(df.columns)
    ws.append([])
    ws.append(headers)

    header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )

    ws.row_dimensions[4].height = 28
    for col_idx, col_name in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

    for row_idx, row in df.iterrows():
        excel_row_idx = row_idx + 5
        row_values = list(row)
        ws.append(row_values)
        ws.row_dimensions[excel_row_idx].height = 22

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=excel_row_idx, column=col_idx)
            cell.border = thin_border
            cell.font = Font(name="Calibri", size=10)

            col_header = headers[col_idx - 1]
            if col_header in ("Referencia", "Modalidad", "Código Sistema"):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_header in ("Precio Actual ($)", "Precio Nuevo ($)", "Costo Diario Actual ($)", "Costo Diario Nuevo ($)", "Diferencia Día ($)"):
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.number_format = '$#,##0.00'
            elif col_header in ("Deprec Actual", "Deprec Nuevo"):
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.number_format = '0.0000'
            elif col_header == "Variación Día (%)":
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.number_format = '+0.0%;-0.0%;0.0%'
            elif col_header == "Estado Sugerido":
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

            # Colorear badges
            if col_header == "Modalidad":
                if cell.value == "PROPIO":
                    cell.fill = PatternFill(start_color="EFF6FF", end_color="EFF6FF", fill_type="solid")
                    cell.font = Font(name="Calibri", size=10, bold=True, color="1D4ED8")
                else:
                    cell.fill = PatternFill(start_color="FFFBEB", end_color="FFFBEB", fill_type="solid")
                    cell.font = Font(name="Calibri", size=10, bold=True, color="B45309")

            if col_header == "Estado Sugerido":
                if "50%" in str(cell.value):
                    cell.fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
                    cell.font = Font(name="Calibri", size=10, bold=True, color="991B1B")
                elif "Moderada" in str(cell.value):
                    cell.fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
                    cell.font = Font(name="Calibri", size=10, bold=True, color="92400E")
                elif "Normal" in str(cell.value):
                    cell.fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
                    cell.font = Font(name="Calibri", size=10, bold=True, color="166534")

    # Ajuste de ancho de columnas
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = max(len(str(cell.value or '')) for cell in col[3:])
        col_name = headers[col[0].column - 1] if col[0].column <= len(headers) else ''
        header_len = len(col_name)
        adjusted_width = max(max_len + 3, header_len + 4, 12)
        if col_letter == "B":
            adjusted_width = min(adjusted_width, 60)
        if col_letter == "L":
            adjusted_width = min(adjusted_width, 45)
        ws.column_dimensions[col_letter].width = adjusted_width

    wb.save(output_path)
    logger.info(f"Excel de equipos guardado en: {output_path}")
    return output_path


def apply_equipments_to_db(
    db: Session,
    results: List[Dict[str, Any]],
) -> int:
    """
    Aplica las actualizaciones de equipos a la base de datos PostgreSQL.
    """
    if not results:
        return 0

    updated = 0
    try:
        for r in results:
            cod = r["codigo_interno"]
            p_nuevo = float(r["precio_nuevo"])
            d_nuevo = float(r["deprec_nuevo"])
            eq = db.query(CostEquipment).filter(CostEquipment.CodEqu == cod).first()
            if eq and p_nuevo > 0:
                eq.precio = p_nuevo
                eq.deprec_factor = d_nuevo
                eq.CosDia = round(p_nuevo * d_nuevo, 4)
                updated += 1
        db.commit()
        logger.info(f"Se actualizaron {updated} equipos en la base de datos.")
        return updated
    except Exception as exc:
        db.rollback()
        logger.error(f"Error aplicando equipos a BD: {exc}", exc_info=True)
        raise


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Auditor y actualizador inteligente de Equipos y Maquinaria para CostBase usando DeepSeek V3 API."
    )
    parser.add_argument(
        "--limit",
        type=int,
        default=None,
        help="Cantidad maxima de equipos a procesar (ej. 25 para prueba piloto, o vacio para todos).",
    )
    parser.add_argument(
        "--batch-size",
        type=int,
        default=25,
        help="Tamano del lote por llamada a DeepSeek. Default: 25.",
    )
    parser.add_argument(
        "--filter",
        type=str,
        default=None,
        help="Filtro de busqueda por descripcion o codigo (ej. 'PALA', 'RETROEXCAVADORA', 'COMPRESOR').",
    )
    parser.add_argument(
        "--deprec-filter",
        choices=["equal_1", "less_1", "alquiler", "propio"],
        default=None,
        help="Filtrar por modalidad: 'equal_1'/'alquiler' (deprec = 1.0) o 'less_1'/'propio' (deprec < 1.0).",
    )
    parser.add_argument(
        "--only-zero-price",
        action="store_true",
        help="Procesar unicamente equipos sin precio asignado ($0 o nulo).",
    )
    parser.add_argument(
        "--output",
        type=str,
        default=None,
        help="Ruta del archivo Excel de salida.",
    )
    parser.add_argument(
        "--apply-to-db",
        action="store_true",
        help="Si se especifica, aplica los nuevos precios y depreciaciones directamente a PostgreSQL tras generar el Excel.",
    )

    args = parser.parse_args()

    if args.batch_size <= 0:
        raise ValueError("El tamano de lote debe ser mayor a 0.")

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_file = args.output or f"precios_equipos_deepseek_{timestamp}.xlsx"

    print("=" * 65)
    print("   COSTBASE — AUDITOR INTELIGENTE DE EQUIPOS CON DEEPSEEK")
    print("=" * 65)
    print(f"Limite maximo       : {args.limit if args.limit else 'TODOS LOS EQUIPOS'}")
    print(f"Tamano de lote      : {args.batch_size}")
    print(f"Filtro de texto     : {args.filter or 'Ninguno'}")
    print(f"Filtro depreciacion : {args.deprec_filter or 'Todas'}")
    print(f"Solo sin precio ($0): {args.only_zero_price}")
    print(f"Archivo de salida   : {output_file}")
    print(f"Actualizar en BD    : {'SI (MODO ESCRITURA)' if args.apply_to_db else 'NO (SOLO AUDITORIA EXCEL)'}")
    print("-" * 65)

    with SessionLocal() as db:
        api_key, base_url, model_name = get_deepseek_credentials(db)
        print(f"Conectado a DeepSeek API: Modelo '{model_name}' ({base_url})")

        print("Consultando equipos en la base de datos...")
        items = fetch_equipments_to_update(
            db=db,
            limit=args.limit,
            search_filter=args.filter,
            deprec_filter=args.deprec_filter,
            only_zero_price=args.only_zero_price,
        )

        total_items = len(items)
        if total_items == 0:
            print("No se encontraron equipos que coincidan con los criterios.")
            sys.exit(0)

        print(f"Total de equipos seleccionados: {total_items}")
        total_batches = (total_items + args.batch_size - 1) // args.batch_size
        print(f"Total de llamadas estimadas a DeepSeek: {total_batches}")
        print("-" * 65)

        start_time = time.time()
        all_results: List[Dict[str, Any]] = []

        for batch_idx in range(total_batches):
            start_i = batch_idx * args.batch_size
            end_i = min(start_i + args.batch_size, total_items)
            current_batch = items[start_i:end_i]

            batch_num = batch_idx + 1
            print(f"[{batch_num}/{total_batches}] Auditando equipos {start_i + 1} al {end_i} con DeepSeek...", end="", flush=True)

            t0 = time.time()
            enriched = call_deepseek_equipment_batch(
                api_key=api_key,
                base_url=base_url,
                model_name=model_name,
                batch_items=current_batch,
            )
            elapsed = time.time() - t0
            all_results.extend(enriched)
            print(f" Completado en {elapsed:.2f}s")

            if batch_idx < total_batches - 1:
                time.sleep(0.5)

        total_time = time.time() - start_time
        print("-" * 65)
        print(f"Procesamiento con IA completado en {total_time:.2f} segundos.")

        print(f"Generando reporte Excel en '{output_file}'...")
        export_equipments_to_excel(all_results, output_file)

        if args.apply_to_db:
            print("Aplicando actualizaciones a la base de datos PostgreSQL...")
            actualizados = apply_equipments_to_db(db, all_results)
            print(f"Base de datos actualizada: {actualizados} equipos guardados.")

        est_in = total_batches * 750
        est_out = total_batches * 900
        cost_in = (est_in / 1_000_000) * 0.27
        cost_out = (est_out / 1_000_000) * 1.10
        total_cost = cost_in + cost_out

        print("=" * 65)
        print("                  RESUMEN DE EJECUCION")
        print("=" * 65)
        print(f"Equipos procesados     : {len(all_results)}")
        print(f"Tiempo total           : {total_time:.1f} s")
        print(f"Llamadas a la API      : {total_batches}")
        print(f"Costo estimado DeepSeek: ${total_cost:.4f} USD")
        print(f"Archivo Excel listo    : {os.path.abspath(output_file)}")
        print("=" * 65)


if __name__ == "__main__":
    main()
