import argparse
import json
import logging
import os
import sys
import time
from datetime import datetime
from typing import Any, Dict, List, Optional, Tuple

from dotenv import load_dotenv
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
import pandas as pd
import requests
from sqlalchemy import or_
from sqlalchemy.orm import Session

# Permitir ejecucion directa agregando la raiz del backend al sys.path
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
BACKEND_DIR = os.path.abspath(os.path.join(SCRIPT_DIR, ".."))
if BACKEND_DIR not in sys.path:
    sys.path.insert(0, BACKEND_DIR)

# Cargar variables de entorno del backend de forma determinista
load_dotenv(os.path.join(BACKEND_DIR, ".env"))

from app.crud.llm import decrypt_api_key
from app.db.models.costbase import CostEquipment, CostLabor, CostMaterial
from app.db.models.llm_provider import LLMProvider
from app.db.session import SessionLocal

logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s")
logger = logging.getLogger("actualizar_precios_deepseek")


def get_deepseek_credentials(db: Session) -> Tuple[str, str, str]:
    """
    Obtiene y desencripta las credenciales de DeepSeek desde la base de datos.
    Retorna (api_key, base_url, model_name).
    """
    if db is None:
        raise ValueError("La sesion de base de datos no puede ser nula.")

    provider = (
        db.query(LLMProvider)
        .filter(LLMProvider.provider_key == "deepseek", LLMProvider.is_active == True)
        .first()
    )
    if not provider:
        raise ValueError("No se encontro un proveedor activo para 'deepseek' en la tabla llm_providers.")

    api_key = decrypt_api_key(provider.api_key_enc)
    if not api_key:
        raise ValueError("La clave API de DeepSeek no pudo ser desencriptada o esta vacia.")

    base_url = "https://api.deepseek.com"
    if provider.extra_params:
        extra = provider.extra_params
        if isinstance(extra, str):
            try:
                extra = json.loads(extra)
            except Exception:
                extra = {}
        if isinstance(extra, dict) and extra.get("base_url"):
            base_url = extra["base_url"].rstrip("/")

    model_name = provider.model_name or "deepseek-chat"
    return api_key, base_url, model_name


def fetch_resources_to_update(
    db: Session,
    resource_type: str,
    limit: Optional[int] = None,
    search_filter: Optional[str] = None,
    only_zero_price: bool = False,
) -> List[Dict[str, Any]]:
    """
    Consulta los insumos a actualizar segun los filtros especificados.
    """
    if not resource_type:
        raise ValueError("El tipo de recurso es obligatorio.")

    items: List[Dict[str, Any]] = []

    if resource_type == "materials":
        query = db.query(CostMaterial)
        if only_zero_price:
            query = query.filter(or_(CostMaterial.CosMat.is_(None), CostMaterial.CosMat == 0.0))
        if search_filter:
            term = f"%{search_filter.strip()}%"
            query = query.filter(
                or_(
                    CostMaterial.CodMat.ilike(term),
                    CostMaterial.ref_code.ilike(term),
                    CostMaterial.Descri.ilike(term),
                )
            )
        query = query.order_by(CostMaterial.ref_code.asc(), CostMaterial.CodMat.asc())
        if limit and limit > 0:
            query = query.limit(limit)

        for m in query.all():
            items.append({
                "referencia": m.ref_code or m.CodMat,
                "codigo_interno": m.CodMat,
                "descripcion": m.Descri or "",
                "unidad": m.UniMat or "UND",
                "precio_actual": float(m.CosMat or 0.0),
            })

    elif resource_type == "equipments":
        query = db.query(CostEquipment)
        if only_zero_price:
            query = query.filter(or_(CostEquipment.precio.is_(None), CostEquipment.precio == 0.0))
        if search_filter:
            term = f"%{search_filter.strip()}%"
            query = query.filter(
                or_(
                    CostEquipment.CodEqu.ilike(term),
                    CostEquipment.ref_code.ilike(term),
                    CostEquipment.Descri.ilike(term),
                )
            )
        query = query.order_by(CostEquipment.ref_code.asc(), CostEquipment.CodEqu.asc())
        if limit and limit > 0:
            query = query.limit(limit)

        for eq in query.all():
            items.append({
                "referencia": eq.ref_code or eq.CodEqu,
                "codigo_interno": eq.CodEqu,
                "descripcion": eq.Descri or "",
                "unidad": "DÍA",
                "precio_actual": float(eq.precio or 0.0),
                "deprec_factor": float(eq.deprec_factor or 1.0),
            })

    elif resource_type == "labors":
        query = db.query(CostLabor)
        if only_zero_price:
            query = query.filter(or_(CostLabor.Jornal.is_(None), CostLabor.Jornal == 0.0))
        if search_filter:
            term = f"%{search_filter.strip()}%"
            query = query.filter(
                or_(
                    CostLabor.CodMan.ilike(term),
                    CostLabor.ref_code.ilike(term),
                    CostLabor.Descri.ilike(term),
                )
            )
        query = query.order_by(CostLabor.ref_code.asc(), CostLabor.CodMan.asc())
        if limit and limit > 0:
            query = query.limit(limit)

        for l in query.all():
            items.append({
                "referencia": l.ref_code or l.CodMan,
                "codigo_interno": l.CodMan,
                "descripcion": l.Descri or "",
                "unidad": "DÍA",
                "precio_actual": float(l.Jornal or 0.0),
                "bono_actual": float(l.Bono or 0.0),
            })

    else:
        raise ValueError(f"Tipo de recurso invalido: '{resource_type}'. Opciones: materials, equipments, labors.")

    return items


def call_deepseek_batch(
    api_key: str,
    base_url: str,
    model_name: str,
    batch_items: List[Dict[str, Any]],
    resource_type: str,
    max_retries: int = 3,
) -> List[Dict[str, Any]]:
    """
    Envia un lote de insumos a DeepSeek para estimacion de precios de mercado.
    Retorna la lista enriquecida con precio_nuevo, variacion_pct y criterio.
    """
    if not batch_items:
        return []

    system_prompt = (
        "Eres un Ingeniero Civil cotizador y Director de Costos de Construccion con vasta "
        "experiencia en precios de mercado de obras civiles, edificaciones y urbanismo en Venezuela "
        "y la region (cotizados en dolares estadounidenses USD referenciales).\n"
        "Tu mision es auditar y actualizar los precios de una lista de insumos de construccion.\n"
        "REGLAS OBLIGATORIAS:\n"
        "1. UNIDAD DE MEDIDA: Respeta estrictamente la unidad indicada (kg, ton, saco 42.5kg, m3, galon, "
        "cunete 5 gal, pieza, tubo 6m, etc.). Un saco de cemento NO cuesta lo mismo que una tonelada.\n"
        "2. HERRAMIENTAS MENORES / MANUALES: Si el recurso es una pala, pico, carretilla, mandarria o "
        "herramienta manual, indica el PRECIO DE COMPRA NUEVO de reposicion de activo (no alquiler diario).\n"
        "3. PRECIOS REALISTAS: Asigna valores de mercado promedio transparentes de distribuidores y ferreterias.\n"
        "4. FORMATO: Responde EXCLUSIVAMENTE con un JSON con el campo 'items', donde cada elemento tenga: "
        "'referencia', 'precio_nuevo' (float en USD), 'variacion_pct' (float con el % de cambio respecto "
        "al actual) y 'criterio' (breve justificacion o especificacion tecnica).\n"
    )

    clean_batch_for_prompt = [
        {
            "referencia": item["referencia"],
            "descripcion": item["descripcion"],
            "unidad": item["unidad"],
            "precio_actual_usd": item["precio_actual"],
        }
        for item in batch_items
    ]

    user_prompt = (
        f"Tipo de insumo: {resource_type.upper()}\n"
        f"Analiza los siguientes {len(clean_batch_for_prompt)} insumos y determina su precio de mercado "
        f"nuevo en USD:\n\n"
        f"{json.dumps(clean_batch_for_prompt, ensure_ascii=False, indent=2)}\n\n"
        f"Devuelve un JSON valido con la estructura:\n"
        f'{{"items": [{{"referencia": "...", "precio_nuevo": 0.0, "variacion_pct": 0.0, "criterio": "..."}}]}}'
    )

    payload = {
        "model": model_name,
        "messages": [
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": user_prompt},
        ],
        "temperature": 0.2,
        "max_tokens": 4096,
        "response_format": {"type": "json_object"},
    }

    headers = {
        "Authorization": f"Bearer {api_key}",
        "Content-Type": "application/json",
    }

    endpoint = f"{base_url}/chat/completions"

    for attempt in range(1, max_retries + 1):
        try:
            resp = requests.post(endpoint, headers=headers, json=payload, timeout=60)
            if resp.status_code == 429:
                wait_time = attempt * 4
                logger.warning(f"Rate limit de DeepSeek alcanzado. Esperando {wait_time}s (intento {attempt}/{max_retries})...")
                time.sleep(wait_time)
                continue

            resp.raise_for_status()
            data = resp.json()
            content = data["choices"][0]["message"]["content"]
            parsed = json.loads(content)
            llm_items = parsed.get("items", [])

            # Indexar respuestas por referencia
            ai_map = {str(it.get("referencia", "")).strip().upper(): it for it in llm_items if isinstance(it, dict)}

            enriched_results: List[Dict[str, Any]] = []
            for item in batch_items:
                ref_key = str(item["referencia"]).strip().upper()
                ai_data = ai_map.get(ref_key, {})

                precio_nuevo = float(ai_data.get("precio_nuevo") or item["precio_actual"] or 0.0)
                precio_actual = float(item["precio_actual"] or 0.0)

                if precio_actual > 0:
                    var_pct = round(((precio_nuevo - precio_actual) / precio_actual) * 100.0, 2)
                else:
                    var_pct = 100.0 if precio_nuevo > 0 else 0.0

                criterio = str(ai_data.get("criterio") or "Estimacion DeepSeek V3 de mercado")

                enriched_results.append({
                    "referencia": item["referencia"],
                    "codigo_interno": item["codigo_interno"],
                    "descripcion": item["descripcion"],
                    "unidad": item["unidad"],
                    "precio_actual": precio_actual,
                    "precio_nuevo": round(precio_nuevo, 4),
                    "diferencia": round(precio_nuevo - precio_actual, 4),
                    "variacion_pct": var_pct,
                    "criterio": criterio,
                })

            return enriched_results

        except Exception as e:
            logger.error(f"Error llamando a DeepSeek en intento {attempt}/{max_retries}: {e}", exc_info=True)
            if attempt == max_retries:
                # En caso de fallo total del lote, preservar precios actuales para no perder datos
                fallback_results = []
                for item in batch_items:
                    fallback_results.append({
                        "referencia": item["referencia"],
                        "codigo_interno": item["codigo_interno"],
                        "descripcion": item["descripcion"],
                        "unidad": item["unidad"],
                        "precio_actual": item["precio_actual"],
                        "precio_nuevo": item["precio_actual"],
                        "diferencia": 0.0,
                        "variacion_pct": 0.0,
                        "criterio": f"Fallo API DeepSeek: {e}",
                    })
                return fallback_results
            time.sleep(attempt * 2)

    return []


def export_to_excel(
    results: List[Dict[str, Any]],
    output_path: str,
    resource_type: str,
) -> str:
    """
    Genera un archivo Excel profesional con las columnas requeridas y formato de celdas.
    """
    if not results:
        raise ValueError("No hay resultados para exportar a Excel.")

    df_data = []
    for r in results:
        # Estado sugerido segun variacion
        var = abs(r["variacion_pct"])
        if r["precio_actual"] == 0:
            estado = "Precio Inicial Fijado"
        elif var > 50:
            estado = "Revisar (Variación > 50%)"
        elif var > 25:
            estado = "Alerta (Variación Moderada)"
        else:
            estado = "Aprobado Normal"

        df_data.append({
            "Referencia": r["referencia"],
            "Descripción": r["descripcion"],
            "Unidad": r["unidad"],
            "Precio Actual ($)": r["precio_actual"],
            "Precio Nuevo ($)": r["precio_nuevo"],
            "Diferencia ($)": r["diferencia"],
            "Variación (%)": r["variacion_pct"] / 100.0,  # En formato porcentaje
            "Criterio / Justificación Mercado": r["criterio"],
            "Estado Sugerido": estado,
            "Código Sistema": r["codigo_interno"],
        })

    df = pd.DataFrame(df_data)

    # Crear libro con openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Precios {resource_type.capitalize()}"

    # Titulo principal
    ws.merge_cells("A1:J1")
    title_cell = ws["A1"]
    title_cell.value = f"COSTBASE PLATFORM — AUDITORÍA Y ACTUALIZACIÓN DE PRECIOS ({resource_type.upper()})"
    title_cell.font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    title_cell.fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 35

    # Subtitulo con metadatos
    ws.merge_cells("A2:J2")
    sub_cell = ws["A2"]
    sub_cell.value = (
        f"Generado con DeepSeek V3 API | Fecha: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')} | "
        f"Total Insumos: {len(results)}"
    )
    sub_cell.font = Font(name="Calibri", size=10, italic=True, color="475569")
    sub_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 20

    # Encabezados de columnas
    headers = list(df.columns)
    ws.append([])  # Fila 3 en blanco
    ws.append(headers)  # Fila 4 encabezados

    header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )

    ws.row_dimensions[4].height = 28
    for col_idx, col_name in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

    # Escribir filas de datos
    for row_idx, row in df.iterrows():
        excel_row_idx = row_idx + 5
        row_values = list(row)
        ws.append(row_values)
        ws.row_dimensions[excel_row_idx].height = 22

        # Formato por celda
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=excel_row_idx, column=col_idx)
            cell.border = thin_border
            cell.font = Font(name="Calibri", size=10)

            # Alineaciones y formatos numericos
            col_header = headers[col_idx - 1]
            if col_header in ("Referencia", "Unidad", "Código Sistema"):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_header in ("Precio Actual ($)", "Precio Nuevo ($)", "Diferencia ($)"):
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.number_format = '$#,##0.00'
            elif col_header == "Variación (%)":
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.number_format = '+0.0%;-0.0%;0.0%'
            elif col_header == "Estado Sugerido":
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

            # Coloreado suave de estado / variacion
            if col_header == "Estado Sugerido":
                if "50%" in cell.value:
                    cell.fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")  # Rojo suave
                    cell.font = Font(name="Calibri", size=10, bold=True, color="991B1B")
                elif "Moderada" in cell.value:
                    cell.fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")  # Amarillo suave
                    cell.font = Font(name="Calibri", size=10, bold=True, color="92400E")
                elif "Normal" in cell.value:
                    cell.fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")  # Verde suave
                    cell.font = Font(name="Calibri", size=10, bold=True, color="166534")

    # Ajustar ancho de columnas automaticamente
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = max(len(str(cell.value or '')) for cell in col[3:])
        col_name = headers[col[0].column - 1] if col[0].column <= len(headers) else ''
        header_len = len(col_name)
        adjusted_width = max(max_len + 3, header_len + 4, 12)
        if col_letter == "B":  # Descripcion
            adjusted_width = min(adjusted_width, 60)
        if col_letter == "H":  # Criterio
            adjusted_width = min(adjusted_width, 45)
        ws.column_dimensions[col_letter].width = adjusted_width

    # Guardar archivo
    wb.save(output_path)
    logger.info(f"Archivo Excel generado exitosamente en: {output_path}")
    return output_path


def apply_updates_to_db(
    db: Session,
    results: List[Dict[str, Any]],
    resource_type: str,
) -> int:
    """
    Aplica directamente los precios nuevos calculados a la base de datos PostgreSQL.
    Solo se ejecuta si se paso la bandera --apply-to-db.
    """
    if not results:
        return 0

    updated_count = 0
    try:
        if resource_type == "materials":
            for r in results:
                codigo = r["codigo_interno"]
                nuevo_precio = float(r["precio_nuevo"])
                mat = db.query(CostMaterial).filter(CostMaterial.CodMat == codigo).first()
                if mat and nuevo_precio > 0:
                    mat.CosMat = nuevo_precio
                    updated_count += 1

        elif resource_type == "equipments":
            for r in results:
                codigo = r["codigo_interno"]
                nuevo_precio = float(r["precio_nuevo"])
                eq = db.query(CostEquipment).filter(CostEquipment.CodEqu == codigo).first()
                if eq and nuevo_precio > 0:
                    factor = eq.deprec_factor if (eq.deprec_factor and eq.deprec_factor > 0) else 1.0
                    eq.precio = nuevo_precio
                    eq.CosDia = round(nuevo_precio * factor, 4)
                    updated_count += 1

        elif resource_type == "labors":
            for r in results:
                codigo = r["codigo_interno"]
                nuevo_precio = float(r["precio_nuevo"])
                labor = db.query(CostLabor).filter(CostLabor.CodMan == codigo).first()
                if labor and nuevo_precio > 0:
                    labor.Jornal = nuevo_precio
                    updated_count += 1

        db.commit()
        logger.info(f"Se actualizaron {updated_count} registros en la tabla {resource_type}.")
        return updated_count
    except Exception as exc:
        db.rollback()
        logger.error(f"Error aplicando cambios a la BD: {exc}", exc_info=True)
        raise


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Actualizador inteligente de precios para CostBase usando DeepSeek V3 API con salida Excel."
    )
    parser.add_argument(
        "--resource",
        choices=["materials", "equipments", "labors"],
        default="materials",
        help="Tipo de recurso a actualizar (materials, equipments, labors). Default: materials.",
    )
    parser.add_argument(
        "--limit",
        type=int,
        default=None,
        help="Cantidad maxima de insumos a procesar (ej. 25 para prueba piloto, o dejar vacio para todos).",
    )
    parser.add_argument(
        "--batch-size",
        type=int,
        default=25,
        help="Tamano del lote enviado a DeepSeek por llamada. Default: 25.",
    )
    parser.add_argument(
        "--filter",
        type=str,
        default=None,
        help="Filtro de busqueda por descripcion o codigo (ej. 'CEMENTO', 'CABILLA', 'ARENA').",
    )
    parser.add_argument(
        "--only-zero-price",
        action="store_true",
        help="Procesar unicamente insumos sin precio asignado ($0 o nulo).",
    )
    parser.add_argument(
        "--output",
        type=str,
        default=None,
        help="Ruta del archivo Excel de salida (default: precios_actualizados_deepseek_<resource>.xlsx).",
    )
    parser.add_argument(
        "--apply-to-db",
        action="store_true",
        help="Si se especifica, aplica los nuevos precios directamente a la base de datos tras generar el Excel.",
    )

    args = parser.parse_args()

    # Validacion temprana
    if args.batch_size <= 0:
        raise ValueError("El tamano de lote (--batch-size) debe ser mayor a 0.")

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_file = args.output or f"precios_actualizados_deepseek_{args.resource}_{timestamp}.xlsx"

    print("=" * 65)
    print("   COSTBASE — ACTUALIZADOR INTELIGENTE DE PRECIOS CON DEEPSEEK")
    print("=" * 65)
    print(f"Recurso objetivo    : {args.resource.upper()}")
    print(f"Limite maximo       : {args.limit if args.limit else 'TODOS'}")
    print(f"Tamano de lote      : {args.batch_size}")
    print(f"Filtro de texto     : {args.filter or 'Ninguno'}")
    print(f"Solo sin precio ($0): {args.only_zero_price}")
    print(f"Archivo de salida   : {output_file}")
    print(f"Actualizar en BD    : {'SI (MODO ESCRITURA)' if args.apply_to_db else 'NO (SOLO AUDITORIA EXCEL)'}")
    print("-" * 65)

    with SessionLocal() as db:
        # 1. Obtener credenciales de DeepSeek
        api_key, base_url, model_name = get_deepseek_credentials(db)
        print(f"Conectado a DeepSeek API: Modelo '{model_name}' ({base_url})")

        # 2. Obtener insumos
        print("Consultando insumos en la base de datos...")
        items = fetch_resources_to_update(
            db=db,
            resource_type=args.resource,
            limit=args.limit,
            search_filter=args.filter,
            only_zero_price=args.only_zero_price,
        )

        total_items = len(items)
        if total_items == 0:
            print("No se encontraron insumos que coincidan con los criterios.")
            sys.exit(0)

        print(f"Total de insumos seleccionados: {total_items}")
        total_batches = (total_items + args.batch_size - 1) // args.batch_size
        print(f"Total de llamadas estimadas a DeepSeek: {total_batches}")
        print("-" * 65)

        start_time = time.time()
        all_results: List[Dict[str, Any]] = []

        # 3. Procesar lotes
        for batch_idx in range(total_batches):
            start_i = batch_idx * args.batch_size
            end_i = min(start_i + args.batch_size, total_items)
            current_batch = items[start_i:end_i]

            batch_num = batch_idx + 1
            print(f"[{batch_num}/{total_batches}] Procesando insumos {start_i + 1} al {end_i} con DeepSeek...", end="", flush=True)

            t0 = time.time()
            enriched = call_deepseek_batch(
                api_key=api_key,
                base_url=base_url,
                model_name=model_name,
                batch_items=current_batch,
                resource_type=args.resource,
            )
            elapsed = time.time() - t0
            all_results.extend(enriched)
            print(f" Completado en {elapsed:.2f}s")

            # Pausa ligera para no golpear rate-limits
            if batch_idx < total_batches - 1:
                time.sleep(0.5)

        total_time = time.time() - start_time
        print("-" * 65)
        print(f"Procesamiento con IA completado en {total_time:.2f} segundos.")

        # 4. Generar archivo Excel
        print(f"Generando reporte Excel en '{output_file}'...")
        export_to_excel(all_results, output_file, args.resource)

        # 5. Aplicar a BD si se solicito
        if args.apply_to_db:
            print("Aplicando actualizaciones aprobadas a la base de datos PostgreSQL...")
            actualizados = apply_updates_to_db(db, all_results, args.resource)
            print(f"Base de datos actualizada: {actualizados} insumos guardados.")

        # Metricas finales y costos
        est_input_tokens = total_batches * 750
        est_output_tokens = total_batches * 900
        cost_in = (est_input_tokens / 1_000_000) * 0.27
        cost_out = (est_output_tokens / 1_000_000) * 1.10
        total_cost = cost_in + cost_out

        print("=" * 65)
        print("                  RESUMEN DE EJECUCION")
        print("=" * 65)
        print(f"Insumos procesados     : {len(all_results)}")
        print(f"Tiempo total           : {total_time:.1f} s")
        print(f"Llamadas a la API      : {total_batches}")
        print(f"Costo estimado DeepSeek: ${total_cost:.4f} USD")
        print(f"Archivo Excel listo    : {os.path.abspath(output_file)}")
        print("=" * 65)


if __name__ == "__main__":
    main()
