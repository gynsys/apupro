from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# (categoría, subcategoría, descripción, unidad, prioridad, especificación)
M = [
    # PLOMERÍA MODERNA
    ("Plomería moderna", "PPR", "Tubo PPR PN20 20 mm", "m", "Alta", "Agua fría/caliente; termofusión"),
    ("Plomería moderna", "PPR", "Tubo PPR PN20 25 mm", "m", "Alta", "Agua fría/caliente; termofusión"),
    ("Plomería moderna", "PPR", "Tubo PPR PN20 32 mm", "m", "Alta", "Agua fría/caliente; termofusión"),
    ("Plomería moderna", "PPR", "Tubo PPR PN20 40 mm", "m", "Media", "Alimentación principal"),
    ("Plomería moderna", "PPR", "Tubo PPR PN20 50 mm", "m", "Media", "Alimentación principal"),
    ("Plomería moderna", "PPR", "Tubo PPR PN20 63 mm", "m", "Media", "Alimentación principal"),
    ("Plomería moderna", "PPR", "Codo PPR 90° 20 mm", "und", "Alta", "Termofusión"),
    ("Plomería moderna", "PPR", "Codo PPR 90° 25 mm", "und", "Alta", "Termofusión"),
    ("Plomería moderna", "PPR", "Codo PPR 90° 32 mm", "und", "Media", "Termofusión"),
    ("Plomería moderna", "PPR", "Tee PPR 20 mm", "und", "Alta", "Termofusión"),
    ("Plomería moderna", "PPR", "Tee PPR 25 mm", "und", "Alta", "Termofusión"),
    ("Plomería moderna", "PPR", "Tee PPR 32 mm", "und", "Media", "Termofusión"),
    ("Plomería moderna", "PPR", "Unión simple PPR 20 mm", "und", "Alta", "Termofusión"),
    ("Plomería moderna", "PPR", "Unión simple PPR 25 mm", "und", "Alta", "Termofusión"),
    ("Plomería moderna", "PPR", "Adaptador PPR rosca macho 20 mm x 1/2\"", "und", "Alta", "Transición a rosca"),
    ("Plomería moderna", "PPR", "Adaptador PPR rosca hembra 20 mm x 1/2\"", "und", "Alta", "Transición a rosca"),
    ("Plomería moderna", "PPR", "Válvula de bola PPR 20 mm", "und", "Alta", "Cierre por sector"),
    ("Plomería moderna", "PPR", "Válvula de bola PPR 25 mm", "und", "Media", "Cierre por sector"),
    ("Plomería moderna", "PPR", "Llave terminal PPR 20 mm", "und", "Media", "Punto de agua"),
    ("Plomería moderna", "PPR", "Reducción PPR 32 x 25 mm", "und", "Media", "Termofusión"),
    ("Plomería moderna", "PPR", "Abrazadera PPR 25 mm", "und", "Media", "Fijación"),
    ("Plomería moderna", "PEX", "Tubo PEX-a 16 mm", "m", "Alta", "Agua potable; curva continua"),
    ("Plomería moderna", "PEX", "Tubo PEX-a 20 mm", "m", "Media", "Agua potable"),
    ("Plomería moderna", "PEX", "Manifold PEX 2 vías", "und", "Media", "Distribución por colector"),
    ("Plomería moderna", "PEX", "Conexión push PEX 16 mm", "und", "Media", "Unión rápida"),
    ("Plomería moderna", "CPVC", "Tubo CPVC 1/2\"", "m", "Alta", "Agua caliente"),
    ("Plomería moderna", "CPVC", "Codo CPVC 90° 1/2\"", "und", "Alta", "Agua caliente"),
    ("Plomería moderna", "CPVC", "Cemento CPVC", "frasco", "Alta", "Solvente"),
    ("Plomería moderna", "CPVC", "Limpiador CPVC", "frasco", "Media", "Preparación de superficie"),
    ("Plomería moderna", "PEAD/PEAL", "Tubo PEAD 20 mm", "m", "Alta", "Agua/gas; enterrado"),
    ("Plomería moderna", "PEAD/PEAL", "Válvula esférica PEAD 20 mm", "und", "Media", "Corte"),
    ("Plomería moderna", "PEAD/PEAL", "Tubo PEAL 1/2\"", "m", "Alta", "Gas natural/GLP"),
    ("Plomería moderna", "Drenaje", "Geotextil no tejido 200 g/m²", "m2", "Media", "Drenaje y filtración"),
    ("Plomería moderna", "Drenaje", "Geodren drenaje lineal", "m2", "Media", "Muros y losas"),
    ("Plomería moderna", "Drenaje", "Sifón PVC doble 2\"", "und", "Media", "Desagüe"),
    ("Plomería moderna", "Drenaje", "Desagüe lineal acero inoxidable", "und", "Media", "Ducha/baño"),
    ("Plomería moderna", "Drenaje", "Tubo PVC sanitario 4\"", "m", "Media", "Drenaje"),
    ("Plomería moderna", "Drenaje", "Rejilla PVC 4\"", "und", "Media", "Drenaje"),
    ("Plomería moderna", "Accesorios", "Abrazadera isofónica 2\"", "und", "Media", "Aislamiento acústico"),
    ("Plomería moderna", "Accesorios", "Válvula reductora de presión 1/2\"", "und", "Media", "Protección de red"),
    ("Plomería moderna", "Accesorios", "Filtro de agua 10\"", "und", "Media", "Potabilización"),
    ("Plomería moderna", "Accesorios", "Contador de agua 1/2\"", "und", "Media", "Medición"),

    # PINTURAS E IMPERMEABILIZACIÓN
    ("Pinturas e impermeabilización", "Elastoméricas", "Pintura elastomérica base agua 5 años", "gln", "Alta", "Techo/muro; alto desempeño"),
    ("Pinturas e impermeabilización", "Elastoméricas", "Pintura elastomérica fibrada 10 años", "gln", "Alta", "Refuerzo de fisuras"),
    ("Pinturas e impermeabilización", "Elastoméricas", "Pintura elastomérica blanca para techo", "gln", "Alta", "Reflectividad solar"),
    ("Pinturas e impermeabilización", "Elastoméricas", "Pintura elastomérica gris para techo", "gln", "Media", "Reflectividad"),
    ("Pinturas e impermeabilización", "Impermeabilizantes", "Membrana líquida de poliuretano", "gln", "Alta", "Impermeabilización"),
    ("Pinturas e impermeabilización", "Impermeabilizantes", "Cemento cristalizante", "kg", "Alta", "Concreto; presión negativa"),
    ("Pinturas e impermeabilización", "Impermeabilizantes", "Manta asfáltica autoadhesiva 3 mm", "rollo", "Alta", "Techos y fundaciones"),
    ("Pinturas e impermeabilización", "Impermeabilizantes", "Manta asfáltica APP 4 mm", "rollo", "Media", "Techos"),
    ("Pinturas e impermeabilización", "Impermeabilizantes", "Imprimante asfáltico", "gln", "Media", "Adherencia"),
    ("Pinturas e impermeabilización", "Impermeabilizantes", "Sellador de poliuretano", "gln", "Alta", "Juntas y fisuras"),
    ("Pinturas e impermeabilización", "Impermeabilizantes", "Espuma de poliuretano expansiva", "und", "Alta", "Relleno y aislamiento"),
    ("Pinturas e impermeabilización", "Impermeabilizantes", "Cinta butílica de aluminio", "rollo", "Alta", "Sellado de juntas"),
    ("Pinturas e impermeabilización", "Impermeabilizantes", "Membrana acrílica para losa", "gln", "Media", "Impermeabilización"),
    ("Pinturas e impermeabilización", "Industriales", "Pintura epóxica para piso", "gln", "Alta", "Tráfico industrial"),
    ("Pinturas e impermeabilización", "Industriales", "Pintura epóxica para pared", "gln", "Media", "Química"),
    ("Pinturas e impermeabilización", "Industriales", "Poliuretano alifático", "gln", "Alta", "Acabado exterior"),
    ("Pinturas e impermeabilización", "Industriales", "Anticorrosivo con zinc", "gln", "Alta", "Estructuras metálicas"),
    ("Pinturas e impermeabilización", "Industriales", "Pintura intumescente", "gln", "Media", "Protección contra fuego"),
    ("Pinturas e impermeabilización", "Industriales", "Pintura antihongos", "gln", "Alta", "Baños/cocinas"),
    ("Pinturas e impermeabilización", "Industriales", "Pintura fotocatalítica", "gln", "Media", "Calidad de aire"),
    ("Pinturas e impermeabilización", "Industriales", "Pintura termoreflectiva", "gln", "Media", "Techos y fachadas"),
    ("Pinturas e impermeabilización", "Industriales", "Pintura para tráfico", "gln", "Media", "Señalización"),
    ("Pinturas e impermeabilización", "Industriales", "Esmalte industrial", "gln", "Media", "Metal/madera"),
    ("Pinturas e impermeabilización", "Industriales", "Primer epóxico", "gln", "Media", "Base de adherencia"),
    ("Pinturas e impermeabilización", "Industriales", "Masilla epóxica", "kg", "Media", "Reparación"),
    ("Pinturas e impermeabilización", "Accesorios", "Diluyente epóxico", "gln", "Media", "Preparación"),

    # PISOS Y REVESTIMIENTOS
    ("Pisos y revestimientos", "Porcelanato", "Porcelanato 60 x 120 cm rectificado mate", "m2", "Alta", "Gran formato"),
    ("Pisos y revestimientos", "Porcelanato", "Porcelanato 60 x 120 cm pulido", "m2", "Alta", "Gran formato"),
    ("Pisos y revestimientos", "Porcelanato", "Porcelanato 120 x 120 cm rectificado", "m2", "Alta", "Gran formato"),
    ("Pisos y revestimientos", "Porcelanato", "Porcelanato 120 x 240 cm", "m2", "Media", "Gran formato"),
    ("Pisos y revestimientos", "Porcelanato", "Porcelanato 90 x 180 cm", "m2", "Media", "Gran formato"),
    ("Pisos y revestimientos", "Porcelanato", "Porcelanato 60 x 60 cm antideslizante", "m2", "Alta", "Exterior"),
    ("Pisos y revestimientos", "Porcelanato", "Porcelanato 30 x 60 cm", "m2", "Media", "Pared"),
    ("Pisos y revestimientos", "Porcelanato", "Porcelanato delgado 6 mm", "m2", "Media", "Remodelación"),
    ("Pisos y revestimientos", "Porcelanato", "Porcelanato símil mármol", "m2", "Alta", "Acabado"),
    ("Pisos y revestimientos", "Porcelanato", "Porcelanato símil madera", "m2", "Alta", "Acabado"),
    ("Pisos y revestimientos", "Porcelanato", "Peldaño de porcelanato", "und", "Media", "Escaleras"),
    ("Pisos y revestimientos", "Porcelanato", "Zócalo de porcelanato", "m", "Media", "Remate"),
    ("Pisos y revestimientos", "Adhesivos", "Adhesivo para porcelanato gran formato", "saco", "Alta", "C2TE"),
    ("Pisos y revestimientos", "Adhesivos", "Adhesivo para porcelanato exterior", "saco", "Alta", "Intemperie"),
    ("Pisos y revestimientos", "Adhesivos", "Boquilla epóxica", "kg", "Alta", "Juntas"),
    ("Pisos y revestimientos", "Adhesivos", "Nivelador de pisos", "kit", "Media", "Instalación"),
    ("Pisos y revestimientos", "Vinílicos", "Piso SPC 4 mm", "m2", "Alta", "Clic; resistencia"),
    ("Pisos y revestimientos", "Vinílicos", "Piso SPC 5 mm", "m2", "Media", "Clic; alto tránsito"),
    ("Pisos y revestimientos", "Vinílicos", "Underlayment acústico", "m2", "Media", "Bajo SPC"),
    ("Pisos y revestimientos", "Vinílicos", "Laminado AC4", "m2", "Media", "Residencial"),
    ("Pisos y revestimientos", "Vinílicos", "Vinílico en rollo", "m2", "Media", "Hospitalario"),
    ("Pisos y revestimientos", "Exteriores", "Deck WPC", "m2", "Alta", "Exterior"),
    ("Pisos y revestimientos", "Exteriores", "Césped artificial", "m2", "Media", "Paisajismo"),
    ("Pisos y revestimientos", "Exteriores", "Microcemento", "m2", "Alta", "Acabado continuo"),
    ("Pisos y revestimientos", "Exteriores", "Piedra flexible", "m2", "Media", "Revestimiento"),
    ("Pisos y revestimientos", "Exteriores", "Panel 3D", "m2", "Media", "Pared interior"),

    # MUROS, TECHOS Y TABIQUERÍA
    ("Muros, techos y tabiquería", "Wall panel", "Wall panel PVC", "m2", "Alta", "Interior"),
    ("Muros, techos y tabiquería", "Wall panel", "Wall panel WPC", "m2", "Alta", "Interior/exterior"),
    ("Muros, techos y tabiquería", "Wall panel", "Wall panel mármol PVC", "m2", "Alta", "Acabado"),
    ("Muros, techos y tabiquería", "Wall panel", "Listón WPC", "m2", "Media", "Fachada"),
    ("Muros, techos y tabiquería", "Wall panel", "Panel acústico", "m2", "Media", "Oficinas"),
    ("Muros, techos y tabiquería", "Drywall", "Plancha drywall 1/2\"", "m2", "Alta", "Tabique"),
    ("Muros, techos y tabiquería", "Drywall", "Plancha drywall RH", "m2", "Alta", "Baño/cocina"),
    ("Muros, techos y tabiquería", "Drywall", "Plancha fibrocemento", "m2", "Media", "Exterior"),
    ("Muros, techos y tabiquería", "Drywall", "Perfil galvanizado 2\"", "m", "Alta", "Estructura"),
    ("Muros, techos y tabiquería", "Drywall", "Perfil galvanizado 3\"", "m", "Media", "Estructura"),
    ("Muros, techos y tabiquería", "Drywall", "Montante galvanizado", "m", "Alta", "Tabique"),
    ("Muros, techos y tabiquería", "Drywall", "Tornillo drywall", "caja", "Alta", "Fijación"),
    ("Muros, techos y tabiquería", "Drywall", "Cinta malla drywall", "rollo", "Alta", "Juntas"),
    ("Muros, techos y tabiquería", "Drywall", "Masilla drywall", "saco", "Alta", "Acabado"),
    ("Muros, techos y tabiquería", "Techos", "Panel sándwich 30 mm", "m2", "Alta", "Cubierta"),
    ("Muros, techos y tabiquería", "Techos", "Teja termoacústica", "m2", "Alta", "Cubierta"),
    ("Muros, techos y tabiquería", "Techos", "Policarbonato alveolar", "m2", "Media", "Cubierta translúcida"),
    ("Muros, techos y tabiquería", "Techos", "Teja translúcida", "m2", "Media", "Iluminación natural"),
    ("Muros, techos y tabiquería", "Techos", "Canaleta de cubierta", "m", "Media", "Drenaje"),
    ("Muros, techos y tabiquería", "Techos", "Cumbrera de cubierta", "m", "Media", "Remate"),
    ("Muros, techos y tabiquería", "Modular", "Panel SIP", "m2", "Media", "Construcción modular"),
    ("Muros, techos y tabiquería", "Modular", "Steel framing", "m2", "Media", "Estructura liviana"),
    ("Muros, techos y tabiquería", "Modular", "Losacero", "m2", "Media", "Losa colaborante"),
    ("Muros, techos y tabiquería", "Refuerzos", "Malla electrosoldada", "m2", "Media", "Refuerzo"),
    ("Muros, techos y tabiquería", "Refuerzos", "Fibra de carbono", "m2", "Media", "Refuerzo estructural"),
    ("Muros, techos y tabiquería", "Refuerzos", "Anclaje químico", "und", "Media", "Fijación estructural"),
    ("Muros, techos y tabiquería", "Refuerzos", "Adhesivo epóxico estructural", "kg", "Media", "Anclaje"),
    ("Muros, techos y tabiquería", "Aislamiento", "Lana mineral", "m2", "Alta", "Térmico/acústico"),
    ("Muros, techos y tabiquería", "Aislamiento", "Poliestireno XPS", "m2", "Alta", "Térmico"),
    ("Muros, techos y tabiquería", "Aislamiento", "Poliestireno EPS", "m2", "Media", "Térmico"),
    ("Muros, techos y tabiquería", "Aislamiento", "Barrera de vapor", "m2", "Media", "Humedad"),
    ("Muros, techos y tabiquería", "Aislamiento", "Aislante reflectivo", "m2", "Media", "Térmico"),

    # ELÉCTRICO E ILUMINACIÓN
    ("Eléctrico e iluminación", "LED", "Tira LED 12 V 60 LED/m", "m", "Alta", "Iluminación indirecta"),
    ("Eléctrico e iluminación", "LED", "Tira LED 24 V 120 LED/m", "m", "Alta", "Alta densidad"),
    ("Eléctrico e iluminación", "LED", "Tira LED RGB", "m", "Media", "Decorativa"),
    ("Eléctrico e iluminación", "LED", "Neón flex LED", "m", "Alta", "Avisos y perfiles"),
    ("Eléctrico e iluminación", "LED", "Perfil de aluminio para LED", "m", "Alta", "Disipación"),
    ("Eléctrico e iluminación", "LED", "Difusor para perfil LED", "m", "Media", "Acabado"),
    ("Eléctrico e iluminación", "LED", "Driver LED 12 V 60 W", "und", "Alta", "Fuente"),
    ("Eléctrico e iluminación", "LED", "Driver LED 24 V 100 W", "und", "Media", "Fuente"),
    ("Eléctrico e iluminación", "LED", "Dimmer LED", "und", "Media", "Control"),
    ("Eléctrico e iluminación", "LED", "Sensor de movimiento", "und", "Alta", "Ahorro"),
    ("Eléctrico e iluminación", "LED", "Panel LED 60 x 60 cm", "und", "Alta", "Oficina"),
    ("Eléctrico e iluminación", "LED", "Downlight LED 6\"", "und", "Alta", "Empotrado"),
    ("Eléctrico e iluminación", "LED", "Spot LED", "und", "Media", "Acento"),
    ("Eléctrico e iluminación", "LED", "Riel electrificado", "m", "Media", "Iluminación flexible"),
    ("Eléctrico e iluminación", "LED", "Proyector LED 100 W", "und", "Alta", "Exterior"),
    ("Eléctrico e iluminación", "LED", "Luminaria solar", "und", "Media", "Exterior"),
    ("Eléctrico e iluminación", "Cables", "Cable THHN 12 AWG", "m", "Alta", "Fuerza"),
    ("Eléctrico e iluminación", "Cables", "Cable THHN 10 AWG", "m", "Alta", "Fuerza"),
    ("Eléctrico e iluminación", "Cables", "Cable THHN 8 AWG", "m", "Media", "Acometida"),
    ("Eléctrico e iluminación", "Cables", "Cable encauchetado 3 x 12", "m", "Alta", "Subterráneo"),
    ("Eléctrico e iluminación", "Protecciones", "Breaker 1P 20 A", "und", "Alta", "Tablero"),
    ("Eléctrico e iluminación", "Protecciones", "Breaker 2P 40 A", "und", "Alta", "Tablero"),
    ("Eléctrico e iluminación", "Protecciones", "Diferencial 2P 30 mA", "und", "Alta", "Protección personal"),
    ("Eléctrico e iluminación", "Protecciones", "Protector contra sobretensiones SPD", "und", "Media", "Tablero"),
    ("Eléctrico e iluminación", "Protecciones", "Tablero 12 circuitos", "und", "Alta", "Distribución"),
    ("Eléctrico e iluminación", "Canalizaciones", "Tubería EMT 1/2\"", "m", "Alta", "Instalación"),
    ("Eléctrico e iluminación", "Canalizaciones", "Tubería IMC 3/4\"", "m", "Media", "Industrial"),
    ("Eléctrico e iluminación", "Canalizaciones", "Caja octogonal", "und", "Alta", "Conexión"),
    ("Eléctrico e iluminación", "Accesorios", "Tomacorriente GFCI", "und", "Alta", "Baño/cocina"),
    ("Eléctrico e iluminación", "Accesorios", "Tomacorriente USB", "und", "Media", "Carga"),
    ("Eléctrico e iluminación", "Accesorios", "Interruptor smart", "und", "Alta", "Domótica"),
    ("Eléctrico e iluminación", "Accesorios", "Contactor 2P 40 A", "und", "Media", "Control"),
    ("Eléctrico e iluminación", "Accesorios", "Temporizador digital", "und", "Media", "Automatización"),
    ("Eléctrico e iluminación", "Accesorios", "UPS 1 kVA", "und", "Media", "Respaldo"),

    # CLIMA Y GAS
    ("Clima y gas", "Climatización", "Minisplit inverter 12.000 BTU", "und", "Alta", "Inverter"),
    ("Clima y gas", "Climatización", "Minisplit inverter 18.000 BTU", "und", "Alta", "Inverter"),
    ("Clima y gas", "Climatización", "Minisplit inverter 24.000 BTU", "und", "Alta", "Inverter"),
    ("Clima y gas", "Climatización", "Unidad VRF exterior", "und", "Media", "Edificios"),
    ("Clima y gas", "Climatización", "Fancoil VRF", "und", "Media", "Interior"),
    ("Clima y gas", "Climatización", "Ducto de fibra", "m2", "Media", "Distribución"),
    ("Clima y gas", "Climatización", "Rejilla de ventilación", "und", "Alta", "Acabado"),
    ("Clima y gas", "Climatización", "Difusor de aire", "und", "Media", "Distribución"),
    ("Clima y gas", "Climatización", "Tubería de cobre 1/4\"", "m", "Alta", "Refrigerante"),
    ("Clima y gas", "Climatización", "Tubería de cobre 3/8\"", "m", "Alta", "Refrigerante"),
    ("Clima y gas", "Climatización", "Aislamiento para tubería", "m", "Alta", "Térmico"),
    ("Clima y gas", "Climatización", "Bomba de condensado", "und", "Media", "Drenaje"),
    ("Clima y gas", "Climatización", "Extractor inline", "und", "Media", "Ventilación"),
    ("Clima y gas", "Climatización", "Filtro HEPA", "und", "Media", "Calidad de aire"),
    ("Clima y gas", "Gas", "Tubería PEAL 1/2\"", "m", "Alta", "Gas"),
    ("Clima y gas", "Gas", "Regulador de gas", "und", "Alta", "GLP"),
    ("Clima y gas", "Gas", "Detector de gas", "und", "Alta", "Seguridad"),
    ("Clima y gas", "Gas", "Válvula de corte de gas", "und", "Alta", "Seguridad"),

    # ENERGÍA SOLAR, DOMÓTICA Y SEGURIDAD
    ("Energía solar, domótica y seguridad", "Solar", "Panel solar 600 W", "und", "Alta", "Fotovoltaico"),
    ("Energía solar, domótica y seguridad", "Solar", "Inversor híbrido 5 kW", "und", "Alta", "Respaldo"),
    ("Energía solar, domótica y seguridad", "Solar", "Batería LiFePO4 5 kWh", "und", "Alta", "Almacenamiento"),
    ("Energía solar, domótica y seguridad", "Solar", "Controlador MPPT", "und", "Alta", "Carga"),
    ("Energía solar, domótica y seguridad", "Solar", "Estructura para panel solar", "und", "Media", "Montaje"),
    ("Energía solar, domótica y seguridad", "Solar", "Cable solar", "m", "Media", "DC"),
    ("Energía solar, domótica y seguridad", "Solar", "Calentador solar", "und", "Media", "ACS"),
    ("Energía solar, domótica y seguridad", "Domótica", "Hub de domótica", "und", "Alta", "Control central"),
    ("Energía solar, domótica y seguridad", "Domótica", "Cerradura smart", "und", "Alta", "Acceso"),
    ("Energía solar, domótica y seguridad", "Domótica", "Cortina motorizada", "und", "Media", "Automatización"),
    ("Energía solar, domótica y seguridad", "Domótica", "Videoportero", "und", "Alta", "Acceso"),
    ("Energía solar, domótica y seguridad", "Seguridad", "Cámara IP", "und", "Alta", "Vigilancia"),
    ("Energía solar, domótica y seguridad", "Seguridad", "NVR", "und", "Alta", "Grabación"),
    ("Energía solar, domótica y seguridad", "Seguridad", "Sensor de alarma", "und", "Alta", "Intrusión"),
    ("Energía solar, domótica y seguridad", "Seguridad", "Detector de humo", "und", "Alta", "Incendio"),
    ("Energía solar, domótica y seguridad", "Seguridad", "Detector de CO", "und", "Media", "Monóxido"),
    ("Energía solar, domótica y seguridad", "Seguridad", "Cargador para vehículo eléctrico", "und", "Media", "EV"),

    # BAÑOS, COCINAS Y CARPINTERÍA
    ("Baños, cocinas y carpintería", "Mesones", "Mesón de cuarzo", "m2", "Alta", "Cocina"),
    ("Baños, cocinas y carpintería", "Mesones", "Mesón de granito", "m2", "Media", "Cocina"),
    ("Baños, cocinas y carpintería", "Mesones", "Mesón solid surface", "m2", "Media", "Baño"),
    ("Baños, cocinas y carpintería", "Baños", "Lavamanos de sobreponer", "und", "Alta", "Baño"),
    ("Baños, cocinas y carpintería", "Baños", "Grifería monomando", "und", "Alta", "Baño"),
    ("Baños, cocinas y carpintería", "Baños", "Sanitario dual", "und", "Alta", "Ahorro de agua"),
    ("Baños, cocinas y carpintería", "Baños", "Mueble de baño", "und", "Alta", "Baño"),
    ("Baños, cocinas y carpintería", "Baños", "Desagüe lineal", "und", "Media", "Ducha"),
    ("Baños, cocinas y carpintería", "Baños", "Regadera termostática", "und", "Media", "Ducha"),
    ("Baños, cocinas y carpintería", "Baños", "Espejo LED", "und", "Media", "Baño"),
    ("Baños, cocinas y carpintería", "Baños", "Mampara de vidrio templado", "m2", "Alta", "Ducha"),
    ("Baños, cocinas y carpintería", "Cocina", "Campana extractora", "und", "Alta", "Cocina"),
    ("Baños, cocinas y carpintería", "Cocina", "Calentador instantáneo", "und", "Media", "ACS"),
    ("Baños, cocinas y carpintería", "Cocina", "Termo eléctrico", "und", "Media", "ACS"),
    ("Baños, cocinas y carpintería", "Carpintería", "Puerta de aluminio", "m2", "Media", "Exterior"),
    ("Baños, cocinas y carpintería", "Carpintería", "Ventana de PVC", "m2", "Media", "Exterior"),
    ("Baños, cocinas y carpintería", "Carpintería", "Mosquitero", "m2", "Media", "Ventana"),
    ("Baños, cocinas y carpintería", "Carpintería", "Persiana", "m2", "Media", "Ventana"),

    # ADHESIVOS, QUÍMICOS Y CONSUMIBLES
    ("Adhesivos, químicos y consumibles", "Selladores", "Silicona neutra", "und", "Alta", "Sellado"),
    ("Adhesivos, químicos y consumibles", "Selladores", "Silicona sanitaria", "und", "Alta", "Baño"),
    ("Adhesivos, químicos y consumibles", "Selladores", "Adhesivo de poliuretano", "und", "Alta", "Construcción"),
    ("Adhesivos, químicos y consumibles", "Selladores", "Adhesivo de contacto", "gln", "Media", "Laminados"),
    ("Adhesivos, químicos y consumibles", "Selladores", "Cinta doble faz", "rollo", "Alta", "Fijación"),
    ("Adhesivos, químicos y consumibles", "Químicos", "Primer universal", "gln", "Alta", "Adherencia"),
    ("Adhesivos, químicos y consumibles", "Químicos", "Thinner", "gln", "Media", "Dilución"),
    ("Adhesivos, químicos y consumibles", "Químicos", "Removedor de pintura", "gln", "Media", "Limpieza"),
    ("Adhesivos, químicos y consumibles", "Químicos", "Catalizador", "und", "Media", "Pinturas"),
    ("Adhesivos, químicos y consumibles", "Químicos", "Resina epóxica", "kg", "Alta", "Reparación"),
    ("Adhesivos, químicos y consumibles", "Químicos", "Lechada epóxica", "kg", "Alta", "Juntas"),
    ("Adhesivos, químicos y consumibles", "Químicos", "Mortero autonivelante", "saco", "Alta", "Nivelación"),
    ("Adhesivos, químicos y consumibles", "Químicos", "Concreto impreso", "m2", "Media", "Pisos decorativos"),
    ("Adhesivos, químicos y consumibles", "Químicos", "Aditivo impermeabilizante", "gln", "Media", "Mortero"),
    ("Adhesivos, químicos y consumibles", "Químicos", "Fibra de polipropileno", "kg", "Media", "Refuerzo de concreto"),
    ("Adhesivos, químicos y consumibles", "Químicos", "Curador de concreto", "gln", "Media", "Curado"),
]


def main():
    out_dir = Path(__file__).resolve().parent
    wb = Workbook()
    ws = wb.active
    ws.title = "Materiales"

    headers = ["Código propuesto", "Descripción", "Unidad", "Categoría", "Prioridad", "Especificación / uso"]
    ws.append(headers)

    header_fill = PatternFill("solid", fgColor="1E3A8A")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    thin = Side(style="thin", color="CBD5E1")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

    prefix = {
        "Plomería moderna": "PLO",
        "Pinturas e impermeabilización": "PIN",
        "Pisos y revestimientos": "PIS",
        "Muros, techos y tabiquería": "MUR",
        "Eléctrico e iluminación": "ELE",
        "Clima y gas": "CLI",
        "Energía solar, domótica y seguridad": "ENS",
        "Baños, cocinas y carpintería": "BCO",
        "Adhesivos, químicos y consumibles": "QUI",
    }
    counters = Counter()
    for cat, sub, desc, unit, prio, spec in M:
        counters[cat] += 1
        code = f"MOD-{prefix.get(cat, 'MAT')}-{counters[cat]:03d}"
        ws.append([code, desc, unit, f"{cat} / {sub}", prio, spec])

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        row[0].font = Font(name="Consolas", size=9)
        if row[4].value == "Alta":
            row[4].fill = PatternFill("solid", fgColor="DCFCE7")
        elif row[4].value == "Media":
            row[4].fill = PatternFill("solid", fgColor="FEF3C7")

    widths = [18, 58, 10, 42, 10, 45]
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:F{ws.max_row}"
    ws.row_dimensions[1].height = 24

    # Resumen
    summary = wb.create_sheet("Resumen")
    summary.append(["Categoría", "Cantidad", "Prioridad alta", "Prioridad media"])
    for cell in summary[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
        cell.border = border

    for cat in prefix:
        cat_rows = [r for r in M if r[0] == cat]
        summary.append([
            cat,
            len(cat_rows),
            sum(1 for r in cat_rows if r[4] == "Alta"),
            sum(1 for r in cat_rows if r[4] == "Media"),
        ])

    for row in summary.iter_rows(min_row=2, max_row=summary.max_row):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="center", wrap_text=True)
    for i, width in enumerate([42, 12, 16, 16], start=1):
        summary.column_dimensions[get_column_letter(i)].width = width
    summary.freeze_panes = "A2"

    notes = wb.create_sheet("Notas")
    notes["A1"] = "Lista propuesta de materiales modernos para CostBase"
    notes["A1"].font = Font(bold=True, size=14, color="1E3A8A")
    notes["A3"] = "Base actual:"
    notes["B3"] = "7.957 materiales"
    notes["A4"] = "Total de materiales propuestos:"
    notes["B4"] = len(M)
    notes["A5"] = "Códigos:"
    notes["B5"] = "Son tentativos (MOD-...). Deben validarse contra la codificación existente."
    notes["A6"] = "Duplicados:"
    notes["B6"] = "Antes de importar, buscar por descripción y referencia para evitar duplicados."
    notes["A7"] = "Precios:"
    notes["B7"] = "La lista no incluye precios; deben cargarse desde el mercado o la base maestra."
    notes["A8"] = "Prioridad alta:"
    notes["B8"] = "Materiales de mayor demanda o adopción reciente; se recomienda cargar primero."
    for cell in ("A3", "A4", "A5", "A6", "A7", "A8"):
        notes[cell].font = Font(bold=True)
    notes.column_dimensions["A"].width = 24
    notes.column_dimensions["B"].width = 90

    path = out_dir / "Lista_Materiales_Modernos_CostBase.xlsx"
    wb.save(path)
    print(path)
    print(f"total={len(M)}")


if __name__ == "__main__":
    main()
