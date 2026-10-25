from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

out = Path(__file__).resolve().parent / "Refuerzo_Fibra_Carbono_CFRP.xlsx"
wb = Workbook()
ws = wb.active
ws.title = "Cotizacion"

header_fill = PatternFill("solid", fgColor="1E3A8A")
header_font = Font(color="FFFFFF", bold=True)
thin = Side(style="thin", color="CBD5E1")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

headers = ["Código BD", "Material", "Unidad", "Precio unitario USD", "Cotización original", "Consumo / aplicación", "Total 100 m² USD", "Observación"]
ws.append(headers)

rows = [
    ("MOD-MUR-033", "Tela de fibra de carbono HM-60 para columnas", "m2", 88.14, "$ 8.814,00 / 100 m²", "1 m²/m²", 8814.00, "Recomendado para columnas"),
    ("MOD-MUR-034", "Tela de fibra de carbono HM-30 para vigas", "m2", 66.96, "$ 6.695,95 / 100 m²", "1 m²/m²", 6695.95, "Recomendado para vigas"),
    ("MOD-MUR-035", "Tela de fibra de carbono HM-20 para vigas", "m2", 54.71, "$ 5.470,56 / 100 m²", "1 m²/m²", 5470.56, "Alternativa para vigas"),
    ("MOD-MUR-036", "Mortero nivelador HM-180 CE kit 20 kg", "kit", 1107.68, "$ 1.107,68 / kit", "4 kit/100 m²", 4430.72, "Mínimo 4 kit por 100 m²"),
    ("MOD-MUR-037", "Imprimante epóxico HM-180 C3P kit 30 kg", "kit", 1256.97, "$ 1.256,97 / kit", "No indicado", None, "Consumo pendiente de confirmar"),
    ("MOD-MUR-038", "Resina epóxica para saturación CFRP", "kg", 47.50, "Sin precio del proveedor", "±1 kg/m²", 4750.00, "Precio referencial de mercado"),
]
for r in rows:
    ws.append(list(r))

for cell in ws[1]:
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = border
for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
    for cell in row:
        cell.border = border
        cell.alignment = Alignment(vertical="center", wrap_text=True)
for i, width in enumerate([16, 54, 9, 18, 30, 24, 20, 42], start=1):
    ws.column_dimensions[get_column_letter(i)].width = width
ws.freeze_panes = "A2"

# Cálculo para 100 m2 columnas
calc = wb.create_sheet("Calculo_100m2")
calc.append(["Sistema para columnas", "Cantidad", "Unidad", "Precio unitario", "Total USD"])
calc_rows = [
    ("Tela HM-60 para columnas", 100, "m2", 88.14, 8814.00),
    ("Mortero nivelador HM-180 CE", 4, "kit", 1107.68, 4430.72),
    ("Imprimante HM-180 C3P (supuesto 1 kit)", 1, "kit", 1256.97, 1256.97),
    ("Resina epóxica (±1 kg/m2)", 100, "kg", 47.50, 4750.00),
]
for r in calc_rows:
    calc.append(r)
calc.append(["Total material directo", "", "", "", 19251.69])
calc.append(["Costo por m2", "", "", "", 192.52])
for cell in calc[1]:
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", wrap_text=True)
for row in calc.iter_rows(min_row=2, max_row=calc.max_row):
    for cell in row:
        cell.border = border
        cell.alignment = Alignment(vertical="center")
for i, width in enumerate([42, 12, 10, 18, 16], start=1):
    calc.column_dimensions[get_column_letter(i)].width = width
calc["A9"] = "Nota: el consumo del imprimante no fue indicado; se usa 1 kit como supuesto."
calc["A9"].font = Font(italic=True, color="9A3412")

# Referencias DB
db = wb.create_sheet("Referencias_BD")
db.append(["Código BD", "Descripción", "Unidad", "Precio BD USD"])
db_rows = [
    ("MOD-MUR-033", "Tela de fibra de carbono HM-60 para refuerzo de columnas", "m2", 88.14),
    ("MOD-MUR-034", "Tela de fibra de carbono HM-30 para refuerzo de vigas", "m2", 66.96),
    ("MOD-MUR-035", "Tela de fibra de carbono HM-20 para refuerzo de vigas", "m2", 54.71),
    ("MOD-MUR-036", "Mortero nivelador HM-180 CE kit 20 kg", "kit", 1107.68),
    ("MOD-MUR-037", "Imprimante epóxico HM-180 C3P kit 30 kg", "kit", 1256.97),
    ("MOD-MUR-038", "Resina epóxica para saturación de fibra de carbono CFRP", "kg", 47.50),
]
for r in db_rows:
    db.append(r)
for cell in db[1]:
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", wrap_text=True)
for row in db.iter_rows(min_row=2, max_row=db.max_row):
    for cell in row:
        cell.border = border
        cell.alignment = Alignment(vertical="center", wrap_text=True)
for i, width in enumerate([18, 60, 10, 16], start=1):
    db.column_dimensions[get_column_letter(i)].width = width

notes = wb.create_sheet("Notas")
notes["A1"] = "Refuerzo estructural con fibra de carbono (CFRP)"
notes["A1"].font = Font(bold=True, size=14, color="1E3A8A")
items = [
    ("Fuente", "Cotización de proveedor compartida por Ing. Pablo Milano, 31/07/2026."),
    ("Moneda", "USD."),
    ("Fabricantes", "HM-60, HM-30, HM-20, HM-180 CE y HM-180 C3P."),
    ("Consumo epoxi", "±1 kg/m² indicado por el proveedor."),
    ("Nivelador", "HM-180 CE kit 20 kg; mínimo 4 kit por 100 m²."),
    ("Imprimante", "El proveedor no indicó consumo por m²; confirmar antes de cerrar el APU."),
    ("Epoxi", "El proveedor no dio precio; se usó un referencial de mercado de $47,50/kg."),
    ("Alcance", "La cotización de tela parece incluir solo el material; no incluye mano de obra ni preparación."),
]
r = 3
for k, v in items:
    notes[f"A{r}"] = k
    notes[f"B{r}"] = v
    notes[f"A{r}"].font = Font(bold=True)
    notes[f"A{r}"].alignment = Alignment(vertical="top")
    notes[f"B{r}"].alignment = Alignment(wrap_text=True, vertical="top")
    r += 1
notes.column_dimensions["A"].width = 20
notes.column_dimensions["B"].width = 110

wb.save(out)
print(out)
